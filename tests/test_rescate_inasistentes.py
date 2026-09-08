#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# ==========================================================================
# This code was generated with the assistance of Claude Sonnet 5 (Anthropic).
# The human author reviewed, modified, and integrated the code.
# Author: Simón Tobar — CESFAM Dr. Luis Ferrada Urzúa (APS, SSMC)
# SPDX-License-Identifier: GPL-3.0-or-later
# ==========================================================================
"""Pruebas de modulos/rem_sm_rescate_inasistentes.py (§8 de
docs/SP_P6_poblacion_plan.md). Datos SINTÉTICOS. Fixtures propias (mismo
layout RAYEN/IRIS que tests/test_sp_p6.py, self-contained a propósito).
Correr:
    python tests/test_rescate_inasistentes.py"""

import sys
import tempfile
from datetime import date
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

import programas.poblacion as pob                          # noqa: E402
import modulos.rem_sp_p6_poblacion as p6mod                 # noqa: E402
import modulos.rem_sm_rescate_inasistentes as resc          # noqa: E402
from programas.rem_utils import ArchivoInvalido             # noqa: E402

_TMP = Path(tempfile.mkdtemp(prefix="autorem_rescate_"))


def _quiet(*_a, **_k):
    pass


# -- Fixtures (subset: solo Depresión 18/19, que exige INSTRUMENTOcontieneMEDIC) --
_Q = {18: "18.- ¿ TIENE  DEPRESIÓN ?", 19: "19.- ESTADO"}
_FORM_HDR = [
    "SERVICIO SALUD", "ESTABLECIMIENTO", "TIPO IDENTIFICACION", "NUMERO TIPO IDENTIFICACION",
    "CODIGO FAMILIA", "NUMERO DE FICHA RAYEN", "NUMERO DE FICHA CODIGO ANTIGUO", "PACIENTE",
    "FECHA DE NACIMIENTO", "EDAD PACIENTE", "AÑO APLICACIÓN FORMULARIO", "MES APLICACIÓN FORMULARIO",
    "DÍAS APLICACIÓN FORMULARIO", "PUEBLO ORIGINARIO", "ALERTAS ADMINISTRATIVAS", "NACIONALIDAD",
    "SEXO", "GENERO", "SECTOR INSCRIPCION", "SECTOR CITA", "DIRECCIÓN", "COMUNA", "TELEFONO 1",
    "TELEFONO 2", "PREVISION", "CONVENIO", "SITUACION", "ESTADO", "FUNCIONARIO PASIVADOR",
    "ATEN ID", "FECHA ATENCION", "FECHA FORMULARIO", "FUNCIONARIO", "INSTRUMENTO",
    "ESTABLECIMIENTO INSCRIPCION", "FORMULARIO", "FUNCIONARIOS FORMULARIO",
] + [_Q[n] for n in sorted(_Q)]


def _mk_formulario(filas):
    p = _TMP / "formulario.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 17):
        ws.cell(row=r, column=1, value=f"banner {r}")
    ws.append(_FORM_HDR)
    filas = filas or [{"rut": "00000000-0", "fecha": date(2020, 1, 1)}]
    for f in filas:
        base = {"NUMERO TIPO IDENTIFICACION": f["rut"], "FECHA FORMULARIO": f["fecha"],
                "INSTRUMENTO": f.get("instr", "Medico"), "SEXO": "Mujer",
                "AÑO APLICACIÓN FORMULARIO": 30}
        ws.append([f.get(h, base.get(h)) for h in _FORM_HDR])
    wb.save(p)
    return p


_INS_HDR = ["TIPO IDENTIFICACION", "NUMERO TIPO IDENTIFICACION", "SEXO", "GENERO",
           "FECHA DE NACIMIENTO", "EDAD AÑOS", "NACIONALIDAD", "PUEBLO INDIG",
           "ALERTAS ADMINISTRATIVAS", "SITUACION", "ESTADO", "SECTOR",
           "FECHA PASIVACION", "MOTIVO PASIVACION"]


def _mk_inscritos(filas):
    p = _TMP / "inscritos.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(_INS_HDR)
    for f in filas:
        base = {"TIPO IDENTIFICACION": f.get("tipoid", "RUN"), "NUMERO TIPO IDENTIFICACION": f["rut"],
               "SEXO": f.get("sexo", "Mujer"), "GENERO": f.get("genero", "Femenina"),
               "FECHA DE NACIMIENTO": f.get("fnac"), "EDAD AÑOS": f.get("edad", 30),
               "NACIONALIDAD": f.get("nacionalidad", "Chilena"), "PUEBLO INDIG": f.get("pueblo", "Ninguno"),
               "ALERTAS ADMINISTRATIVAS": f.get("alertas", ""), "SITUACION": f.get("situacion", "Inscrito"),
               "ESTADO": f.get("estado", "Activo"), "SECTOR": f.get("sector", "Norte"),
               "FECHA PASIVACION": f.get("fpasiv"), "MOTIVO PASIVACION": f.get("mpasiv", "")}
        ws.append([base.get(h, "") for h in _INS_HDR])
    wb.save(p)
    return p


_ADA_HDR = ["NUMERO TIPO IDENTIFICACION", "ATEN ID", "FECHA ATENCION", "ACTIVIDADES",
           "DIAGNOSTICOS", "INSTRUMENTO", "TIPO ATENCION", "SEXO", "SECTOR"]


def _mk_ada(filas):
    p = _TMP / "ada.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(_ADA_HDR)
    filas = filas or [{"rut": "00000000-0", "fecha": date(2020, 1, 1), "act": "Consulta general"}]
    for f in filas:
        ws.append([f["rut"], f.get("id", "A"), f["fecha"].strftime("%d/%m/%Y"), f.get("act", ""),
                  f.get("diag", ""), f.get("instr", "Medico"), f.get("tipo", "Consulta"),
                  f.get("sexo", "Mujer"), f.get("sector", "Norte")])
    wb.save(p)
    return p


def _sm(rut, fecha):
    """Atención SM (una de las 7 actividades validadas) -> alimenta Activo12m/rescate."""
    return {"rut": rut, "fecha": fecha, "act": "Controles Salud Mental"}


MES = (2026, 8)   # corte = 2026-08-31


def _poblacion(formulario_filas, inscritos_filas, ada_filas, mes=MES, **kw):
    return pob.construir_poblacion(
        str(_mk_inscritos(inscritos_filas)), str(_mk_formulario(formulario_filas)),
        str(_mk_ada(ada_filas)), mes=mes, log=_quiet, **kw)


# ======================================================================
# Rescate_6m / Rescate_13m
# ======================================================================
def test_rescate_6m_trae_sector_y_ultima_atencion():
    """Última atención SM hace 6 meses exactos (feb-2026 para el corte ago-2026) ->
    cae en Rescate_6m con Sector y Fecha última atención SM correctos."""
    P = _poblacion([], [{"rut": "11111111-1", "sector": "Sur"}],
                   [_sm("11111111-1", date(2026, 2, 10))])
    tablas = resc.construir_rescate(P, pob.cargar_atenciones(_mk_ada([_sm("11111111-1", date(2026, 2, 10))]), log=_quiet), mes=MES, log=_quiet)
    r6 = tablas["Rescate_6m"]
    assert "11111111-1" in set(r6["RUN"])
    fila = r6[r6["RUN"] == "11111111-1"].iloc[0]
    assert fila["Sector"] == "Sur"
    assert fila["Fecha última atención SM"].date() == date(2026, 2, 10)


def test_fallecido_no_se_excluye_pero_se_flagea():
    """§8.3 (corregido sep-2026): Motivo Pasivación=Fallecido NO saca a la persona
    de Rescate_6m — sigue ahí, y además queda flageada en Posibles_Fallecidos
    (mismo tratamiento que Posibles_Traslados)."""
    ada_filas = [_sm("22222222-2", date(2026, 2, 10))]
    P = _poblacion([], [{"rut": "22222222-2", "mpasiv": "Fallecido", "estado": "Pasivo"}], ada_filas)
    d_ada = pob.cargar_atenciones(_mk_ada(ada_filas), log=_quiet)
    tablas = resc.construir_rescate(P, d_ada, mes=MES, log=_quiet)
    assert "22222222-2" in set(tablas["Rescate_6m"]["RUN"])
    assert "22222222-2" in set(tablas["Posibles_Fallecidos"]["RUN"])
    assert tablas["Posibles_Fallecidos"].iloc[0]["Motivo Pasivación"] == "Fallecido"


def test_fallecidos_mes_exige_fecha_pasivacion_en_el_mes_reportado():
    P = _poblacion(
        [{"rut": "33333333-3", "fecha": date(2026, 1, 5), **{_Q[18]: "SI", _Q[19]: "19.- INGRESO"}}],
        [{"rut": "33333333-3", "estado": "Pasivo", "mpasiv": "Fallecido", "fpasiv": date(2026, 8, 15)}],
        [])
    d_ada = pob.cargar_atenciones(_mk_ada([]), log=_quiet)
    tablas = resc.construir_rescate(P, d_ada, mes=MES, log=_quiet)
    assert "33333333-3" in set(tablas["Fallecidos_mes"]["RUN"])


def test_fallecidos_mes_no_incluye_fallecidos_de_otro_mes():
    P = _poblacion(
        [{"rut": "44444444-4", "fecha": date(2026, 1, 5), **{_Q[18]: "SI", _Q[19]: "19.- INGRESO"}}],
        [{"rut": "44444444-4", "estado": "Pasivo", "mpasiv": "Fallecido", "fpasiv": date(2026, 3, 1)}],
        [])
    d_ada = pob.cargar_atenciones(_mk_ada([]), log=_quiet)
    tablas = resc.construir_rescate(P, d_ada, mes=MES, log=_quiet)
    assert "44444444-4" not in set(tablas["Fallecidos_mes"]["RUN"])


def test_posibles_traslados_no_excluye_de_rescate():
    """§8.5: un traslado NO sale de Rescate_6m — aparece en AMBAS tablas."""
    ada_filas = [_sm("55555555-5", date(2026, 2, 10))]
    P = _poblacion([], [{"rut": "55555555-5", "mpasiv": "Traslado a otro centro"}], ada_filas)
    d_ada = pob.cargar_atenciones(_mk_ada(ada_filas), log=_quiet)
    tablas = resc.construir_rescate(P, d_ada, mes=MES, log=_quiet)
    assert "55555555-5" in set(tablas["Rescate_6m"]["RUN"])
    assert "55555555-5" in set(tablas["Posibles_Traslados"]["RUN"])
    assert tablas["Posibles_Traslados"].iloc[0]["Motivo Pasivación"] == "Traslado a otro centro"


# ======================================================================
# Brecha_Medico (§8.6)
# ======================================================================
def test_brecha_medico_detecta_dx_registrado_solo_por_no_medico():
    P_med = _poblacion(
        [{"rut": "66666666-6", "fecha": date(2026, 8, 1), "instr": "Psicólogo(a)",
         **{_Q[18]: "SI", _Q[19]: "19.- INGRESO"}}],
        [{"rut": "66666666-6"}], [_sm("66666666-6", date(2026, 8, 1))])
    assert P_med.loc[P_med["Número"] == "66666666-6", "¿Ingresado?"].iloc[0] == "NO"

    insc = pob.cargar_inscritos(_mk_inscritos([{"rut": "66666666-6"}]), log=_quiet)
    form = pob.cargar_formulario_sm(_mk_formulario(
        [{"rut": "66666666-6", "fecha": date(2026, 8, 1), "instr": "Psicólogo(a)",
         **{_Q[18]: "SI", _Q[19]: "19.- INGRESO"}}]), log=_quiet)
    d_ada = pob.cargar_atenciones(_mk_ada([_sm("66666666-6", date(2026, 8, 1))]), log=_quiet)
    brecha = resc.calcular_brecha_medico(insc, form, d_ada, P_med, mes=MES, log=_quiet)
    assert "66666666-6" in set(brecha["RUN"])
    fila = brecha[brecha["RUN"] == "66666666-6"].iloc[0]
    assert "Depresión" in fila["Dx que entran solo por no-médico"]
    assert fila["Estamento que lo registró"] == "Psicólogo(a)"


def test_brecha_medico_vacia_si_el_dx_ya_es_medico():
    P_med = _poblacion(
        [{"rut": "77777777-7", "fecha": date(2026, 8, 1), "instr": "Medico",
         **{_Q[18]: "SI", _Q[19]: "19.- INGRESO"}}],
        [{"rut": "77777777-7"}], [_sm("77777777-7", date(2026, 8, 1))])
    insc = pob.cargar_inscritos(_mk_inscritos([{"rut": "77777777-7"}]), log=_quiet)
    form = pob.cargar_formulario_sm(_mk_formulario(
        [{"rut": "77777777-7", "fecha": date(2026, 8, 1), "instr": "Medico",
         **{_Q[18]: "SI", _Q[19]: "19.- INGRESO"}}]), log=_quiet)
    d_ada = pob.cargar_atenciones(_mk_ada([_sm("77777777-7", date(2026, 8, 1))]), log=_quiet)
    brecha = resc.calcular_brecha_medico(insc, form, d_ada, P_med, mes=MES, log=_quiet)
    assert "77777777-7" not in set(brecha["RUN"])


# ======================================================================
# Guardarraíl §8.6: el P6 rechaza una población con el filtro médico apagado
# ======================================================================
def test_p6_rechaza_poblacion_con_exigir_medico_apagado():
    P_todos = _poblacion(
        [{"rut": "88888888-8", "fecha": date(2026, 8, 1), "instr": "Psicólogo(a)",
         **{_Q[18]: "SI", _Q[19]: "19.- INGRESO"}}],
        [{"rut": "88888888-8"}], [], exigir_medico=False)
    try:
        p6mod.construir_p6(P_todos, log=_quiet)
        assert False, "debió rechazar un P con exigir_medico=False"
    except ArchivoInvalido:
        pass


def test_procesar_rechaza_P_con_exigir_medico_apagado():
    P_todos = _poblacion([], [{"rut": "99999999-9"}], [], exigir_medico=False)
    try:
        resc.procesar(str(_mk_inscritos([{"rut": "99999999-9"}])),
                      str(_mk_formulario([])), str(_mk_ada([])), mes=MES, log=_quiet, P=P_todos)
        assert False, "debió rechazar un P con exigir_medico=False"
    except ArchivoInvalido:
        pass


# ======================================================================
# Integración: procesar() + escribir() -> LEEME primera, 6 hojas
# ======================================================================
def test_escribir_hoja_leeme_primera_y_las_6_hojas():
    ada_filas = [_sm("10000001-1", date(2026, 2, 10))]
    inscritos = _mk_inscritos([{"rut": "10000001-1", "sector": "Sur"}])
    formulario = _mk_formulario([])
    ada = _mk_ada(ada_filas)
    E = resc.procesar(str(inscritos), str(formulario), str(ada), mes=MES, log=_quiet)
    salida = _TMP / "rescate.xlsx"
    resc.escribir(E, salida)

    wb = openpyxl.load_workbook(salida)
    assert wb.sheetnames[0] == "LEEME"
    for hoja in ("Rescate_6m", "Rescate_13m", "Fallecidos_mes", "Posibles_Fallecidos",
                "Posibles_Traslados", "Brecha_Medico"):
        assert hoja in wb.sheetnames, wb.sheetnames


def _main():
    pruebas = [v for k, v in sorted(globals().items())
              if k.startswith("test_") and callable(v)]
    fallos = 0
    for fn in pruebas:
        try:
            fn(); print(f"PASS  {fn.__name__}")
        except AssertionError as e:
            fallos += 1; print(f"FAIL  {fn.__name__} -> {e or 'assert'}")
        except Exception as e:  # noqa: BLE001
            import traceback
            fallos += 1; print(f"ERROR {fn.__name__} -> {type(e).__name__}: {e}")
            traceback.print_exc()
    print("-" * 50)
    print(f"{len(pruebas) - fallos}/{len(pruebas)} OK" + (f" ({fallos} problemas)" if fallos else ""))
    return 1 if fallos else 0


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    sys.exit(_main())
