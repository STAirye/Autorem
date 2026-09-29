# -*- coding: utf-8 -*-
# This code was generated with the assistance of Claude Sonnet 5.5 (Anthropic).
# The human author reviewed, modified, and integrated the code.
#
# Author: Simon Tobar - CESFAM Dr. Luis Ferrada Urzua (APS, SSMC)
# Copyright (C) 2026 Simon Tobar
# SPDX-License-Identifier: GPL-3.0-or-later
"""Lectura de exports con python-calamine (docs/lectura_calamine_plan.md): tipos
exactos, perdidas aceptadas, paridad con la referencia openpyxl, regla de hoja y no-xlsx."""
import sys
import zipfile
from datetime import date, datetime, time
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _aislar_cache   # noqa: E402,F401  (PRIMERO: nunca tocar el ~/.autorem real)

from programas.rem_utils import (ArchivoInvalido, abrir_xlsx_ro, cargar_canonico,   # noqa: E402
                                 filas_hoja, filas_xlsx, primeras_filas)

RAIZ = Path(__file__).resolve().parent.parent


def _libro(ruta, *hojas):
    """hojas = [(nombre, [fila, ...]), ...]; escrito con openpyxl."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas:
        ws = wb.create_sheet(nombre)
        for f in filas:
            ws.append(f)
    wb.save(ruta)
    return ruta


def test_tipos_exactos(tmp_path):
    fila = [1, 2.5, 10000000, True, datetime(2026, 8, 1), datetime(2026, 8, 1, 13, 5),
            time(13, 5), "0123", "ñandú"]
    got = filas_xlsx(_libro(tmp_path / "t.xlsx", ("H", [fila])))[0]
    esperado = [(int, 1), (float, 2.5), (int, 10000000), (bool, True),
                (datetime, datetime(2026, 8, 1)), (datetime, datetime(2026, 8, 1, 13, 5)),
                (time, time(13, 5)), (str, "0123"), (str, "ñandú")]
    for v, (tipo, val) in zip(got, esperado):
        assert type(v) is tipo and v == val, (v, tipo)


def test_perdidas_aceptadas_quedan_fijas(tmp_path):
    """Pin de un comportamiento CONOCIDO de calamine (no un requisito): celda de error y
    string de solo espacios salen ''. Si un cambio de version lo mueve, que se note."""
    got = filas_xlsx(_libro(tmp_path / "p.xlsx", ("H", [["#N/A", "  ", "x"]])))[0]
    assert got[0] == "" and got[1] == "" and got[2] == "x"


def _referencia(ruta):
    """Filas de la unica hoja con datos, por el camino openpyxl de siempre."""
    wb = abrir_xlsx_ro(ruta)
    try:
        hojas = [filas_hoja(ws) for ws in wb.worksheets]
    finally:
        wb.close()
    return [h for h in hojas if any(v not in (None, "") for f in h for v in f)]


def _sin_cola(filas):
    """`filas` sin las filas completamente vacias del final."""
    filas = list(filas)
    while filas and all(v in (None, "") for v in filas[-1]):
        filas.pop()
    return filas


def test_paridad_con_openpyxl_en_refs_tablas():
    revisados = 0
    for ruta in sorted((RAIZ / "refs_tablas").glob("*.xlsx")):
        ref = _referencia(ruta)
        if len(ref) != 1:
            continue
        # Las filas VACIAS del final no cuentan: openpyxl entrega las que solo traen
        # formato (el catalogo ENO: 57 con dato + 70 vacias) y calamine no. Lo que se
        # exige es que no falte ni cambie ninguna fila CON dato (2.0.20: fallaba solo
        # en el arbol del autor, por un .xlsx local sin versionar).
        nuevo, fref = _sin_cola(filas_xlsx(ruta)), _sin_cola(ref[0])
        assert len(nuevo) == len(fref), f"{ruta.name}: distinto n de filas"
        for i, (fn, fr) in enumerate(zip(nuevo, fref)):
            fr = tuple("" if v is None else v for v in fr)
            fr += ("",) * (len(fn) - len(fr))
            fn += ("",) * (len(fr) - len(fn))
            for j, (a, b) in enumerate(zip(fn, fr)):
                # Perdida conocida (la misma familia que «solo espacios -> ''»): en un
                # texto que el .xlsx no marca `xml:space="preserve"`, calamine recorta los
                # espacios del BORDE (el Maestro real: '  Club de adulto mayor...').
                # `norm()` los recorta igual, asi que ninguna comparacion cambia. Solo se
                # tolera ESO; cualquier otra diferencia falla.
                # (Solo el blanco de XML: el \xa0 de RAYEN NO se recorta, y `str.strip()` si.)
                borde = (isinstance(a, str) and isinstance(b, str) and a != b
                         and a == b.strip(" \t\r\n"))
                assert borde or repr(a) == repr(b), f"{ruta.name} fila {i} col {j}: {a!r} vs {b!r}"
        revisados += 1
    assert revisados, "refs_tablas/ no trae ningun .xlsx comparable"


def test_regla_de_hojas(tmp_path):
    with pytest.raises(ArchivoInvalido) as ex:
        filas_xlsx(_libro(tmp_path / "v.xlsx", ("H", [])))
    assert ex.value.categoria == "sin_datos"
    with pytest.raises(ArchivoInvalido) as ex:
        filas_xlsx(_libro(tmp_path / "d.xlsx", ("A", [["x"]]), ("B", [["y"]])))
    assert ex.value.categoria == "modificado"
    # datos en una hoja que NO es la activa (la activa es la 1a, vacia): se leen
    assert filas_xlsx(_libro(tmp_path / "n.xlsx", ("Vacia", []), ("Datos", [["ok"]]))) == [("ok",)]


def test_no_xlsx(tmp_path):
    from gui import runner
    falso = tmp_path / "Formulario_Rayen.xlsx"
    falso.write_text("<html><body>x</body></html>")
    with pytest.raises(zipfile.BadZipFile) as ex:
        primeras_filas(falso, 5)
    assert runner.es_error_formato(ex.value)
    with pytest.raises(ArchivoInvalido) as ai:
        cargar_canonico(falso, lambda h: {}, ["RUN"])
    assert ai.value.categoria == "no_legible"


def test_la_grilla_arranca_en_A1_aunque_los_datos_empiecen_en_C3(tmp_path):
    """Las dos trampas de `iter_rows` de calamine (2.0.17): se salta las COLUMNAS vacias de
    la izquierda (hay que rellenarlas o todo indice de columna se corre) y hace
    `PanicException` en Rust sobre una hoja vacia (RAYEN siempre trae dos)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["C3"] = "banner"
    ws["C5"], ws["D5"], ws["E5"] = "RUN", "FECHA", 2.5
    ws["C6"], ws["D6"] = "r1", 5
    wb.create_sheet("vacia1"); wb.create_sheet("vacia2")
    p = tmp_path / "c3.xlsx"
    wb.save(p)
    got = filas_xlsx(p)
    ref = [tuple("" if v is None else v for v in f) for f in _referencia(p)[0]]
    assert [tuple(map(repr, f)) for f in got] == [tuple(map(repr, f)) for f in ref]
    assert got[2][2] == "banner" and got[4][2:5] == ("RUN", "FECHA", 2.5)


def test_el_error_de_no_xlsx_no_lleva_la_carpeta(tmp_path):
    """La ruta completa lleva la carpeta de OneDrive con el nombre del usuario, y el texto
    termina en el dialogo y en el log (revision ciega, 2.0.17): solo el nombre."""
    falso = tmp_path / "Formulario_Rayen.xlsx"
    falso.write_text("<html><body>x</body></html>")
    for leer in (filas_xlsx, lambda p: primeras_filas(p, 5)):
        with pytest.raises(zipfile.BadZipFile) as ex:
            leer(falso)
        assert falso.name in str(ex.value) and str(tmp_path) not in str(ex.value)


def test_primeras_filas(tmp_path):
    p = _libro(tmp_path / "f.xlsx", ("H", [[i, "a"] for i in range(20)]))
    assert primeras_filas(p, 5) == filas_xlsx(p)[:5]
    assert len(primeras_filas(p, 5)) == 5
