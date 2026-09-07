#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# ==========================================================================
# This code was generated with the assistance of Claude Sonnet 5 (Anthropic).
# The human author reviewed, modified, and integrated the code.
# Author: Simón Tobar — CESFAM Dr. Luis Ferrada Urzúa (APS, SSMC)
# SPDX-License-Identifier: GPL-3.0-or-later
# ==========================================================================
"""Pruebas de programas/cobertura.py (hoja «LEEME», ver docs/hoja_cobertura_plan.md).
    python tests/test_cobertura.py"""

import importlib
import pkgutil
import sys
import tempfile
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

import modulos                        # noqa: E402
from programas import cobertura       # noqa: E402

_TMP = Path(tempfile.mkdtemp(prefix="autorem_cobertura_"))


def _quiet(*_a, **_k): pass


# -- Test 1: guardarraíl anti-olvido (§7.1) -----------------------------
def _ids_con_salida():
    """ids esperados en COBERTURA: todo módulo de `modulos/` con `escribir()` o
    `TAREA` (las dos formas que hoy usan los módulos para producir un .xlsx),
    derivando el id igual que la convención de nombre (rem_<id> -> <id>)."""
    ids = set()
    for info in pkgutil.iter_modules(modulos.__path__):
        mod = importlib.import_module(f"modulos.{info.name}")
        if hasattr(mod, "escribir") or hasattr(mod, "TAREA"):
            nombre = info.name[4:] if info.name.startswith("rem_") else info.name
            ids.add(nombre)
    return ids


def test_catalogo_cubre_todos_los_modulos():
    esperados = _ids_con_salida()
    assert esperados, "no encontré ningún módulo con escribir()/TAREA en modulos/"
    assert esperados == set(cobertura.COBERTURA.keys()), (
        f"desfase entre modulos/ y COBERTURA: faltan={esperados - set(cobertura.COBERTURA)} "
        f"sobran={set(cobertura.COBERTURA) - esperados}")


# -- Test 2: la hoja es LA PRIMERA, en los dos caminos de escritura (§7.2) --
def test_hoja_primera_camino_pandas():
    # Camino pandas (§6.1): pd.ExcelWriter arranca con el libro VACÍO (pandas saca
    # la hoja default de openpyxl.Workbook() al abrir en modo escritura).
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    cobertura.escribir_hoja(wb, "sm_actividades", {"mes": (2026, 7)})
    ws2 = wb.create_sheet("SM_Detalle")
    ws2.append(["x"])
    assert wb.sheetnames[0] == "LEEME", wb.sheetnames


def test_hoja_primera_camino_a05():
    # Camino openpyxl (§6.2): el libro YA trae la hoja de datos original en el índice 0.
    wb = openpyxl.Workbook()
    wb.active.title = "Datos_Originales"
    wb.active.append(["dato", "no", "tocar"])
    cobertura.escribir_hoja(wb, ["a05_o_egresos", "a05_n_ingresos"],
                            {"mes": (2026, 7), "archivos": ["export.xlsx"]})
    assert wb.sheetnames[0] == "LEEME", wb.sheetnames
    assert "Datos_Originales" in wb.sheetnames
    assert wb["Datos_Originales"]["A1"].value == "dato"   # intacta


# -- Test 3/4: avisos de ESTA corrida (§7.3/7.4) -------------------------
_ADA_HDR = ["NUMERO TIPO IDENTIFICACION", "ATEN ID", "FECHA ATENCION", "ACTIVIDADES",
            "DIAGNOSTICOS", "INSTRUMENTO", "TIPO ATENCION", "SEXO", "AÑOS ATENCION"]
_ADA_K = {"run": 0, "id": 1, "fecha": 2, "act": 3, "dg": 4, "instr": 5, "tipo": 6,
          "sexo": 7, "edad": 8}


def _mk_ada(rows):
    p = _TMP / "ada.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(_ADA_HDR)
    for r in rows:
        line = [""] * len(_ADA_HDR)
        for k, v in r.items():
            line[_ADA_K[k]] = v
        ws.append(line)
    wb.save(p)
    return p


def _leer_textos(ws):
    return [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]


def test_avisos_sin_grupal_aparecen_en_la_hoja():
    import modulos.rem_sm_actividades as smact
    ada = _mk_ada([dict(run="1-9", id="A1", fecha="10/07/2026", act="Consulta De Salud Mental",
                        dg="x", instr="Medico(a)", tipo="Espontanea", sexo="Femenino", edad="30 años")])
    E = smact.procesar(ada, grupal=None, mes=(2026, 7), log=_quiet)
    assert any("A06 psicosocial" in a[0] for a in E.attrs["avisos"]), E.attrs["avisos"]

    salida = _TMP / "sm_sin_grupal.xlsx"
    smact.escribir(E, salida)
    wb = openpyxl.load_workbook(salida)
    assert wb.sheetnames[0] == "LEEME"
    textos = _leer_textos(wb["LEEME"])
    assert any("A06 psicosocial" in t for t in textos), textos


def test_sin_avisos_dice_sin_avisos():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = cobertura.escribir_hoja(wb, "sp_p6_poblacion", {"mes": (2026, 7)}, avisos=())
    textos = _leer_textos(ws)
    assert any("Sin avisos" in t for t in textos), textos


# -- Test 5: todo el texto del catálogo es cp1252-encodable (§7.5) -------
def test_catalogo_cp1252_encodable():
    for modulo_id, cat in cobertura.COBERTURA.items():
        cat["rem"].encode("cp1252")
        for c in cat["cubre"]:
            c.encode("cp1252")
        for fila in cat["no_cubre"]:
            for campo in fila:
                campo.encode("cp1252")


def _main():
    pruebas = [v for k, v in sorted(globals().items())
               if k.startswith("test_") and callable(v)]
    fallos = 0
    for fn in pruebas:
        try:
            fn(); print(f"PASS  {fn.__name__}")
        except AssertionError as e:
            fallos += 1; print(f"FAIL  {fn.__name__} -> {e or 'assert'}")
        except Exception as e:  # noqa: BLE001
            import traceback
            fallos += 1; print(f"ERROR {fn.__name__} -> {type(e).__name__}: {e}")
            traceback.print_exc()
    print("-" * 50)
    print(f"{len(pruebas) - fallos}/{len(pruebas)} OK" + (f" ({fallos} problemas)" if fallos else ""))
    return 1 if fallos else 0


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    sys.exit(_main())
