#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# ==========================================================================
# This code was generated with the assistance of Claude Opus 5 (Anthropic).
# The human author reviewed, modified, and integrated the code.
#
# Author: Simon Tobar - CESFAM Dr. Luis Ferrada Urzua (APS, SSMC)
# Copyright (C) 2026 Simon Tobar
# SPDX-License-Identifier: GPL-3.0-or-later
# This program is free software: you can redistribute it and/or modify it
# under the terms of the GNU General Public License as published by the
# Free Software Foundation, either version 3 of the License, or (at your
# option) any later version. Distributed WITHOUT ANY WARRANTY. See the GNU
# General Public License for more details: <https://www.gnu.org/licenses/>.
# ==========================================================================
"""
Fase 2 del eje de formatos: clasificar la FUENTE del grupo pandas.

Lo que se protege acá es la distincion de TRES estados. Dos serian suficientes
para atajar el Monitoreo Admin, pero dejarian un falso positivo permanente el dia
que RAYEN renombre una columna del IRIS -- y un aviso que grita siempre deja de
leerse, que es como se pierde el fail-loud.
"""

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import _aislar_cache   # noqa: E402,F401  (PRIMERO: nunca tocar el ~/.autorem real)

from programas import formatos
from programas.rem_utils import (leer_xlsx, resolver_columnas, MAPA_ATENCIONES, norm,
                                 cargar_canonico, ArchivoInvalido)

RAIZ = Path(__file__).resolve().parent.parent
ADA_IRIS = RAIZ / "refs_tablas" / "ATENCIONESDIAGNOSTICOSACTIVIDADES_iris.xlsx"
MONITOREO = RAIZ / "refs_tablas" / "Monitoreo_de_Actividades_anonimizado.xlsx"

TODAS = formatos.SOLO_IRIS_ATENCIONES


def _col(presentes):
    """Simula el {canonico: columna|None} de resolver_columnas."""
    return {k: (f"col_{k}" if k in presentes else None) for k in TODAS}


def _con_una_fila(origen, destino):
    """Copia un export de `refs_tablas/` agregandole UNA fila sintetica de relleno.

    Las planillas de ejemplo del repo son HEADER-ONLY a proposito (privacidad), pero
    `cargar_canonico` ahora exige filas de datos (guarda sobre la FUENTE, CLAUDE.md
    regla 2). Los tests que necesitan el ENCABEZADO REAL como guardarrail -y no los
    datos- pasan por aca: el header sigue siendo el del export de verdad."""
    import openpyxl
    wb = openpyxl.load_workbook(origen)
    ws = wb.active
    ws.append(["x"] * ws.max_column)
    wb.save(destino)
    return destino


# -- Los tres estados ---------------------------------------------------------

def test_todas_las_claves_presentes_es_fuente_plena():
    estado, ausentes = formatos.clasificar_fuente(_col(TODAS))
    assert estado == formatos.FUENTE_PLENA
    assert ausentes == []


def test_ninguna_clave_es_fuente_parcial():
    estado, ausentes = formatos.clasificar_fuente(_col([]))
    assert estado == formatos.FUENTE_PARCIAL
    assert set(ausentes) == set(TODAS)


def test_algunas_claves_es_export_cambiado_no_parcial():
    """RAYEN renombra UNA columna del IRIS: el archivo SIGUE siendo el IRIS pleno.
    Marcarlo 'parcial' seria un falso positivo recurrente -> estado propio."""
    estado, ausentes = formatos.clasificar_fuente(_col(TODAS[1:]))
    assert estado == formatos.FUENTE_CAMBIADA
    assert ausentes == [TODAS[0]]


def test_cambiada_aunque_falte_casi_todo_menos_una():
    """El limite inferior de 'cambiada': con UNA sola clave viva no es 'parcial'."""
    estado, _ = formatos.clasificar_fuente(_col([TODAS[0]]))
    assert estado == formatos.FUENTE_CAMBIADA


# -- El aviso que llega a la hoja LEEME ---------------------------------------

def test_fuente_plena_no_genera_aviso():
    assert formatos.aviso_fuente(formatos.FUENTE_PLENA, [], "lo que sea") is None


def test_aviso_parcial_apunta_al_usuario():
    av = formatos.aviso_fuente(formatos.FUENTE_PARCIAL, list(TODAS), "CONSECUENCIA X")
    casilla, estado, motivo, que_hacer = av
    assert estado == "FUENTE PARCIAL"
    assert "CONSECUENCIA X" in motivo
    assert "IRIS" in que_hacer          # la accion es del usuario: bajar el otro export


def test_aviso_cambiado_apunta_al_dev():
    """Distinto destinatario: no es 'cargaste el archivo equivocado', es
    'RAYEN movio el piso'."""
    av = formatos.aviso_fuente(formatos.FUENTE_CAMBIADA, ["ATENID"], "CONSECUENCIA Y")
    _casilla, estado, motivo, que_hacer = av
    assert estado == "EXPORT CAMBIADO"
    assert "ATENID" in motivo
    assert "dev" in que_hacer.lower()


def test_el_aviso_tiene_la_forma_que_espera_la_hoja_leeme():
    av = formatos.aviso_fuente(formatos.FUENTE_PARCIAL, ["ATENID"], "x")
    assert len(av) == 4 and all(isinstance(p, str) for p in av)


# -- Contra el export IRIS REAL (no una maqueta) ------------------------------

@pytest.mark.skipif(not ADA_IRIS.exists(), reason="falta el export IRIS de ejemplo")
def test_el_ada_iris_real_clasifica_como_plena():
    """Guardarrail anti-falso-positivo: el export IRIS versionado tiene que dar
    'plena'. Si este test se cae, la firma se desincronizo del export real."""
    hdr, _filas = leer_xlsx(ADA_IRIS)
    col = resolver_columnas(hdr, MAPA_ATENCIONES)
    estado, ausentes = formatos.clasificar_fuente(col)
    assert estado == formatos.FUENTE_PLENA, f"claves no resueltas: {ausentes}"


@pytest.mark.skipif(not ADA_IRIS.exists(), reason="falta el export IRIS de ejemplo")
def test_cargar_canonico_deja_el_veredicto_en_attrs(tmp_path):
    d, _col = cargar_canonico(_con_una_fila(ADA_IRIS, tmp_path / "ada.xlsx"), None,
                              lambda h: resolver_columnas(h, MAPA_ATENCIONES),
                              solo_iris=formatos.SOLO_IRIS_ATENCIONES,
                              log=lambda *a, **k: None)
    assert d.attrs["fuente"][0] == formatos.FUENTE_PLENA


@pytest.mark.skipif(not MONITOREO.exists(), reason="falta el Monitoreo de ejemplo")
def test_el_monitoreo_real_clasifica_como_parcial():
    """El otro extremo, contra el archivo REAL (header-only, sep-2026): el
    'Monitoreo de Actividades' del eje Administrativo tiene que dar 'parcial', con
    las CINCO claves demograficas ausentes. Es la contraparte del test del ADA IRIS: uno protege
    contra falsos positivos, este contra falsos negativos."""
    hdr, _filas = leer_xlsx(MONITOREO)
    col = resolver_columnas(hdr, MAPA_ATENCIONES)
    estado, ausentes = formatos.clasificar_fuente(col)
    assert estado == formatos.FUENTE_PARCIAL
    assert set(ausentes) == set(TODAS)


@pytest.mark.skipif(not MONITOREO.exists(), reason="falta el Monitoreo de ejemplo")
def test_el_monitoreo_igual_pasa_las_requeridas_de_cargar_atenciones():
    """Por esto el bug era SILENCIOSO y no un ArchivoInvalido: el Monitoreo resuelve
    TODAS las columnas requeridas, asi que cargaba sin chistar. Lo que faltaba no
    tenia guardia. Si algun dia esto empieza a fallar, `requeridas` cambio y hay que
    revisar si el camino admin del A23 sigue vivo."""
    hdr, _filas = leer_xlsx(MONITOREO)
    col = resolver_columnas(hdr, MAPA_ATENCIONES)
    for k in ("RUN", "FECHA", "ACT", "DIAG", "INSTR", "TIPO"):
        assert col.get(k), f"{k} dejo de resolver en el Monitoreo"


# -- Equivalencias admin encontradas con la muestra (v1.9.3) ------------------

@pytest.mark.skipif(not ADA_IRIS.exists(), reason="falta el export IRIS de ejemplo")
def test_en_iris_la_edad_del_rem_nunca_cae_al_anos_pelado():
    """GUARDARRAIL de la trampa semantica: 'AÑOS' significa cosas DISTINTAS en los
    dos reportes (IRIS = edad a la DESCARGA; Monitoreo = edad a la ATENCION). El
    mapa resuelve por orden, asi que en IRIS debe ganar SIEMPRE 'AÑOS ATENCION'.
    Si RAYEN lo renombrara, el fallback daria edad-a-la-descarga en silencio: este
    test es lo unico que lo atrapa."""
    col = resolver_columnas(leer_xlsx(ADA_IRIS)[0], MAPA_ATENCIONES)
    assert "ATENCION" in norm(col["ANOS_AT"]), col["ANOS_AT"]
    assert norm(col["ANOS_AT"]) != norm(col["ANOS"])   # no son la misma columna


@pytest.mark.skipif(not MONITOREO.exists(), reason="falta el Monitoreo de ejemplo")
def test_en_el_monitoreo_las_equivalencias_admin_resuelven():
    """El Monitoreo no tiene ATEN ID ni 'AÑOS ATENCION', pero SI equivalentes:
    'N°' agrupa las filas de una atencion, y su 'AÑOS' ya es a la atencion
    (confirmado por el autor contra enero-2026)."""
    col = resolver_columnas(leer_xlsx(MONITOREO)[0], MAPA_ATENCIONES)
    assert col["ATENID"] == "N°"
    assert norm(col["ANOS_AT"]) == norm("AÑOS")
    assert col["PROF"] == "FUNCIONARIO"


@pytest.mark.skipif(not MONITOREO.exists(), reason="falta el Monitoreo de ejemplo")
def test_atenid_del_monitoreo_no_entra_en_la_firma_iris():
    """Al darle equivalente admin a ATENID hubo que sacarlo de SOLO_IRIS_ATENCIONES:
    si no, el Monitoreo clasificaria 'cambiada' (mensaje para el dev) en vez de
    'parcial' (mensaje para el usuario)."""
    assert "ATENID" not in formatos.SOLO_IRIS_ATENCIONES
    col = resolver_columnas(leer_xlsx(MONITOREO)[0], MAPA_ATENCIONES)
    assert formatos.clasificar_fuente(col)[0] == formatos.FUENTE_PARCIAL


def test_el_correlativo_se_namespacea_por_archivo(tmp_path):
    """El 'N°' del Monitoreo REINICIA en 1 en cada export. Al concatenar dos
    archivos, dos atenciones distintas compartirian id y el conteo por-atencion
    subcontaria EN SILENCIO. Deben quedar 4 ids distintos, no 2."""
    import openpyxl
    from programas.rem_utils import ATENID_CORRELATIVO
    paths = []
    for n in ("ene.xlsx", "feb.xlsx"):
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append([ATENID_CORRELATIVO, "RUN", "FECHA CONSULTA",
                   "ACTIVIDAD Y/O PROCEDIMIENTO", "DIAGNOSTICO", "INSTRUMENTO",
                   "TIPO DE ATENCION"])
        for i in (1, 2):        # el MISMO correlativo en los dos archivos
            ws.append([i, "11111111-1", "01/01/2026", "act", "dx", "Medico", "Esp"])
        p = tmp_path / n; wb.save(p); paths.append(p)

    d, _col = cargar_canonico(list(paths), None,
                              lambda h: resolver_columnas(h, MAPA_ATENCIONES),
                              log=lambda *a, **k: None)
    assert d["ATENID"].nunique() == 4, "dos atenciones distintas se fusionaron"
    assert all("|" in v for v in d["ATENID"])


def test_el_aten_id_de_iris_NO_se_namespacea(tmp_path):
    """Contrapartida: el ATEN ID de IRIS es global y unico, asi que si el mismo
    aparece en dos exports que se solapan tiene que DEDUPLICAR, no contarse dos
    veces. Namespacearlo romperia eso."""
    import openpyxl
    paths = []
    for n in ("a.xlsx", "b.xlsx"):
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["ATEN ID", "NUMERO TIPO IDENTIFICACION", "FECHA ATENCION",
                   "ACTIVIDADES", "DIAGNOSTICOS", "INSTRUMENTO", "TIPO ATENCION"])
        ws.append([9001, "11111111-1", "01/01/2026", "act", "dx", "Medico", "Esp"])
        p = tmp_path / n; wb.save(p); paths.append(p)

    d, _col = cargar_canonico(list(paths), None,
                              lambda h: resolver_columnas(h, MAPA_ATENCIONES),
                              log=lambda *a, **k: None)
    assert d["ATENID"].nunique() == 1, "el ATEN ID global no debe namespacearse"


def test_sin_solo_iris_no_clasifica_nada(tmp_path):
    """Compatibilidad: los cargadores que no pasan `solo_iris` siguen igual."""
    d, _col = cargar_canonico(_con_una_fila(ADA_IRIS, tmp_path / "ada.xlsx"), None,
                              lambda h: resolver_columnas(h, MAPA_ATENCIONES),
                              log=lambda *a, **k: None)
    assert "fuente" not in d.attrs


@pytest.mark.skipif(not ADA_IRIS.exists(), reason="falta el export IRIS de ejemplo")
def test_fuente_parcial_se_loguea_ruidosa(tmp_path):
    """No basta con dejarlo en attrs: tiene que salir por el log (fail-loud)."""
    dicho = []
    # `h[0]` y no "RUN" a secas: la unica clave que resuelve tiene que apuntar a una
    # columna que EXISTA en el header (antes esto pasaba de casualidad, porque con el
    # export header-only no habia ninguna fila que indexar).
    cargar_canonico(_con_una_fila(ADA_IRIS, tmp_path / "ada.xlsx"), None,
                    lambda h: {k: None for k in TODAS} | {"RUN": h[0]},
                    solo_iris=TODAS, log=lambda m: dicho.append(str(m)))
    assert any("PARCIAL" in m for m in dicho)


# -- Cruce ADA <-> grupal (formatos.parece_reporte / verificar_cruce) --

def test_parece_reporte_reconoce_cada_lado():
    ada = ["NUMERO TIPO IDENTIFICACION", "FECHA ATENCION", "ATEN ID",
           "DIAGNOSTICOS", "ACTIVIDADES", "ALERTAS ADMINISTRATIVAS"]
    gru = ["FUNCIONARIO PRESTADOR", "RUN FUNCIONARIO", "ACTIVIDADES",
           "TIPO PARTICIPANTE", "ASISTE"]
    assert formatos.parece_reporte(ada) == "ada"
    assert formatos.parece_reporte(gru) == "grupal"


def test_parece_reporte_sin_evidencia_no_acusa():
    """Empate (incluido 0-0) -> None: nunca se acusa un cruce sin evidencia."""
    assert formatos.parece_reporte(["COLUMNA A", "COLUMNA B"]) is None


def test_verificar_cruce_solo_dispara_al_reves():
    ada = ["ATEN ID", "DIAGNOSTICOS", "ALERTAS ADMINISTRATIVAS"]
    formatos.verificar_cruce(ada, "ada", "x.xlsx")          # el esperado: no levanta
    with pytest.raises(ArchivoInvalido) as ex:
        formatos.verificar_cruce(ada, "grupal", "x.xlsx")   # cruzado: levanta
    assert ex.value.categoria == "cruzados"
    assert "multitudes" in str(ex.value)


# -- detectar_eje sobre FILAS vs sobre la hoja completa -----------------------
A05_IRIS = RAIZ / "refs_tablas" / "Formularios_RAYEN_csm_IRis.xlsx"
A05_ADMIN = RAIZ / "refs_tablas" / "Formulario_csm_reporte_Administrativo.xlsx"


def _cheap(ruta):
    """Lo que hace la GUI al elegir un archivo: leer SOLO el encabezado."""
    from programas.rem_utils import primeras_filas
    return primeras_filas(ruta, formatos.MAX_FILAS_HEADER)


def _full(ruta):
    import openpyxl
    return openpyxl.load_workbook(ruta, data_only=True).active


@pytest.mark.parametrize("ruta,espera", [(A05_IRIS, "iris"), (A05_ADMIN, "administrativo")])
def test_detectar_eje_filas_coincide_con_la_hoja_completa(ruta, espera):
    """Las dos puertas tienen que dar el MISMO veredicto sobre los exports reales.

    La GUI 2.0 detectaba el formato con un `load_workbook` sin `read_only` -- o sea
    parseando el export completo -- y despues el worker lo parseaba otra vez para
    procesarlo: DOS lecturas completas por corrida. `detectar_eje_filas` deja hacerlo
    con las primeras MAX_FILAS_HEADER filas en modo `read_only`; este test es el que
    amarra que ese atajo no cambie la respuesta."""
    assert formatos.detectar_eje_filas(_cheap(ruta)) == espera
    assert formatos.detectar_eje(_full(ruta)) == espera


def test_detectar_eje_filas_no_necesita_el_resto_del_archivo(tmp_path):
    """Las firmas viven en el encabezado: un export con miles de filas se clasifica
    igual viendo solo las primeras MAX_FILAS_HEADER. Si alguien mueve la deteccion a
    una columna de DATOS, esto lo caza."""
    import openpyxl
    p = tmp_path / "admin_grande.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([formatos.ADMIN_BANNER])                  # banner en A1
    for _ in range(7):
        ws.append([None])
    ws.append(["RUT", "Edad de registro formulario", "Fecha formulario", "Sexo"])
    for i in range(3000):
        ws.append([f"1111111{i % 10}-1", "45 anios", "06/07/2026", "Mujer"])
    wb.save(p)

    completo = _cheap(p)
    assert len(completo) == formatos.MAX_FILAS_HEADER, "el corte por islice no se aplico"
    assert formatos.detectar_eje_filas(completo) == "administrativo"
    # Y con MENOS filas todavia (solo el banner + el header) sigue acertando.
    assert formatos.detectar_eje_filas(completo[:9]) == "administrativo"
    assert formatos.detectar_eje(_full(p)) == "administrativo"


# -- La <dimension> del .xlsx NO manda (ronda 5) -------------------------------
# En modo read_only openpyxl acota `iter_rows` a la <dimension> que declara el
# archivo. Con la etiqueta rota se perdian filas/columnas SIN ningun error: el ADA
# contaba de menos, el A05 rechazaba un IRIS valido, y el escaner de PII no veia lo
# de fuera. `rem_utils.abrir_xlsx_ro` la descarta; estos tests la rompen a proposito.

def _con_dimension(origen, destino, ref, hoja="xl/worksheets/sheet1.xml"):
    """Copia un .xlsx reescribiendo la <dimension> de `hoja` a `ref` (o
    insertandola si no la trae)."""
    import zipfile
    with zipfile.ZipFile(origen) as zi, zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename == hoja:
                xml = data.decode("utf-8")
                i = xml.find("<dimension ")
                if i >= 0:
                    xml = xml[:i] + f'<dimension ref="{ref}"/>' + xml[xml.find(">", i) + 1:]
                else:
                    j = xml.find(">", xml.find("<worksheet")) + 1
                    xml = xml[:j] + f'<dimension ref="{ref}"/>' + xml[j:]
                data = xml.encode("utf-8")
            zo.writestr(it, data)
    return destino


def _read_only_pelado(ruta, n):
    """Lo que hacia el codigo ANTES: prueba que el fixture de verdad arma la trampa."""
    from itertools import islice
    import openpyxl
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        return list(islice(wb.active.iter_rows(values_only=True), n))
    finally:
        wb.close()


@pytest.mark.parametrize("ref", ["A1:D5", "A1"])
def test_leer_xlsx_no_trunca_con_la_dimension_rota(tmp_path, ref):
    """El cuello de botella de TODO el grupo pandas (ADA, grupal, NSP, Inscritos...).
    Truncar aca es el peor bug posible: un numero de menos con cara de legitimo."""
    import openpyxl
    src = tmp_path / "diez.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["RUN", "FECHA", "ACT", "DIAG"])
    for i in range(10):
        ws.append([f"r{i}", "01/08/2026", "x", "y"])
    wb.save(src)
    roto = _con_dimension(src, tmp_path / "roto.xlsx", ref)
    assert len(_read_only_pelado(roto, 100)) < 11, "el fixture ya no reproduce la trampa"
    hdr, filas = leer_xlsx(roto)
    assert len(hdr) == 4 and len(filas) == 10, (len(hdr), len(filas))
    assert all(len(f) == 4 for f in filas), "las filas tienen que salir del mismo ancho"


def test_verificar_hoja_unica_ve_datos_fuera_de_la_dimension(tmp_path):
    """Una hoja extra con datos (tabla dinamica agregada) que su <dimension> no declara
    pasaba por vacia, y el export modificado entraba como bueno."""
    import openpyxl
    from programas.rem_utils import verificar_hoja_unica
    src = tmp_path / "dos_hojas.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["RUN", "FECHA"])
    wb.active.append(["r1", "01/08/2026"])
    ws2 = wb.create_sheet("Pivote")
    ws2["C3"] = "suma"
    wb.save(src)
    roto = _con_dimension(src, tmp_path / "dos_hojas_roto.xlsx", "A1",
                          hoja="xl/worksheets/sheet2.xml")
    with pytest.raises(ArchivoInvalido) as ex:
        verificar_hoja_unica(roto)
    assert ex.value.categoria == "modificado"


def test_verificar_hoja_unica_cierra_el_archivo_aunque_la_lectura_reviente(monkeypatch):
    """read_only deja el .xlsx ABIERTO hasta `close()`. Sin `finally`, un export a medio
    sincronizar por OneDrive que revienta a mitad de la lectura quedaba bloqueado
    mientras el dialogo de error seguia abierto."""
    import programas.rem_utils as ru
    cerrado = []

    class HojaRota:
        title = "Hoja1"

        def iter_rows(self, **_kw):
            raise EOFError("zip truncado")

    class LibroFalso:
        worksheets = [HojaRota()]

        def close(self):
            cerrado.append(True)

    monkeypatch.setattr(ru, "abrir_xlsx_ro", lambda _e: LibroFalso())
    with pytest.raises(EOFError):
        ru.verificar_hoja_unica("export.xlsx")
    assert cerrado, "el workbook quedo abierto tras la excepcion"


def test_deteccion_a05_con_dimension_rota(tmp_path):
    """El IRIS real con la <dimension> en `A1`: un read_only pelado ve el encabezado
    con UNA columna y lo daba 'desconocido' (-> 'Formato no reconocido' sobre un
    archivo valido). Regresion que metio la ronda 4 al leer solo el encabezado."""
    roto = _con_dimension(A05_IRIS, tmp_path / "iris_dim_a1.xlsx", "A1")
    assert formatos.detectar_eje_filas(
        _read_only_pelado(roto, formatos.MAX_FILAS_HEADER)) == "desconocido", \
        "el fixture ya no reproduce la trampa"
    assert formatos.detectar_eje_filas(_cheap(roto)) == "iris"


def test_scan_catalogo_ve_ruts_fuera_de_la_dimension(tmp_path):
    """La guarda de PII del About (CLAUDE.md regla 1): un RUT fuera de la <dimension>
    no se escaneaba. El RUT se ARMA aca (DV calculado) para no dejar uno literal en
    el repo -- el hook anti-RUT tambien mira los tests."""
    import openpyxl
    from tools.hook_pre_commit_rut import dv, sospechosos
    from tools.scan_catalogo import escanear
    cuerpo = "1" + "5432876"
    rut = f"{cuerpo}-{dv(cuerpo)}"
    assert sospechosos(rut), "el RUT sintetico tiene que parecerle real al detector"
    src = tmp_path / "catalogo.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(10):
        ws.append([f"C{i:02d}", "descripcion"])
    ws.append(["J00", f"paciente {rut}"])
    wb.save(src)
    roto = _con_dimension(src, tmp_path / "catalogo_roto.xlsx", "A1:B5")
    hallazgos, _, _ = escanear(roto)
    assert [h for h in hallazgos if h[0] == "RUT"], "un RUT fuera de la <dimension> paso el escaneo"


def test_catalogo_deis_se_lee_entero_con_la_dimension_rota(tmp_path):
    """`catalogos._hojas` (lo que carga un catalogo elegido a mano en el About) tiene
    que ver lo MISMO que el escaner de PII que lo aprobo: todo el archivo."""
    import openpyxl
    from programas import catalogos
    src = tmp_path / "cat.xlsx"
    wb = openpyxl.Workbook()
    for i in range(10):
        wb.active.append([f"C{i:02d}", "descripcion"])
    wb.save(src)
    roto = _con_dimension(src, tmp_path / "cat_roto.xlsx", "A1:B5")
    (_titulo, filas), = catalogos._hojas(roto)
    assert len(filas) == 10, len(filas)


# -- Varios archivos: la FUENTE se clasifica por archivo (ronda 5) --------------

def test_fuente_multi_archivo_no_depende_del_orden(tmp_path):
    """[IRIS, Monitoreo] daba 'plena' (se miraba solo el PRIMERO): banner verde 'A/D/A
    de IRIS completo' y sin aviso en LEEME, sobre filas sin demografia. [Monitoreo,
    IRIS] daba 'parcial'. Ahora los dos ordenes dan lo mismo, y el aviso nombra al
    archivo que no es IRIS."""
    import shutil
    from programas.rem_utils import cargar_atenciones
    iris = _con_una_fila(ADA_IRIS, tmp_path / "iris.xlsx")
    moni = _con_una_fila(MONITOREO, tmp_path / "moni.xlsx")
    for orden in ([iris, moni], [moni, iris]):
        d = cargar_atenciones(orden, log=lambda *_: None)
        assert d.attrs["fuente"][0] == formatos.FUENTE_PARCIAL, (orden, d.attrs["fuente"])
        assert d.attrs["fuente_mezcla"] == ["moni.xlsx"]
    av = formatos.aviso_fuente(*d.attrs["fuente"], "CONSECUENCIA",
                               archivos=d.attrs["fuente_mezcla"])
    assert "moni.xlsx" in av[2] and "MENOS" in av[2]

    iris2 = shutil.copy(iris, tmp_path / "iris2.xlsx")
    d = cargar_atenciones([iris, iris2], log=lambda *_: None)
    assert d.attrs["fuente"][0] == formatos.FUENTE_PLENA
    assert d.attrs["fuente_mezcla"] is None
