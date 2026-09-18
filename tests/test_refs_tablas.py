#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# ==========================================================================
# This code was generated with the assistance of Claude Opus 5 (Anthropic).
# The human author reviewed, modified, and integrated the code.
#
# Author: Simón Tobar — CESFAM Dr. Luis Ferrada Urzúa (APS, SSMC)
# SPDX-License-Identifier: GPL-3.0-or-later
# ==========================================================================
"""
Pruebas del pipeline de referencias (refs_tablas/ + tools/limpiar_refs.py).

Hasta la ronda 11 de la revision de gui-2.0 (sep-2026) no habia NINGUNA: el script no
arrancaba en el Python del proyecto (3.9) sin que nadie lo notara, partia los
encabezados de dos filas, y varias referencias versionadas estaban editadas a mano (sin
banner, con 'Columna1' donde el export dice 'ASISTE (SI/NO)'). Datos SINTETICOS.
Correr: python tests/test_refs_tablas.py
"""

import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.worksheet.table import Table

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(REPO / "tools"))
import _aislar_cache   # noqa: E402,F401  (PRIMERO: nunca tocar el ~/.autorem real)

import limpiar_refs as L                    # noqa: E402
from programas.rem_utils import dv_rut      # noqa: E402
from tools.scan_catalogo import escanear    # noqa: E402

_TMP = Path(tempfile.mkdtemp(prefix="autorem_refs_"))
RUT = "11111111-1"
# Un RUT que el escaner SI caza: el 11111111-1 y los cuerpos 1000... son placeholders y
# los deja pasar a proposito. Se arma en runtime (cuerpo aritmetico + DV calculado) para
# no dejar un literal con forma de RUT en el repo (CLAUDE.md §8.1).
_CUERPO = str(2 * 11000007)
RUT_CAZABLE = f"{_CUERPO}-{dv_rut(_CUERPO)}"


def _crudo(nombre, banner_extra=None):
    """Export 'crudo' sintetico con la forma de RAYEN: banner, encabezado de DOS pisos
    (celdas combinadas, como Utilizacion de Cupos), filas de datos con RUT, una tabla
    de Excel y 2 hojas vacias."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws.append(["Comuna", "Maipu"])
    ws.append(["Establecimiento", "[CESFAM] Ejemplo"])
    ws.append(banner_extra or ["Mes", "AGOSTO"])
    ws.append([])
    ws.append(["Profesional", "Instrumento", "Sector", "Rendimiento", "Cupos no utilizados", None, None])
    ws.append([None, None, None, None, "Total", "Bloqueados", "Libres"])
    ws.merge_cells("A5:A6")
    ws.merge_cells("E5:G5")
    for i in range(30):
        ws.append([f"Persona {i}", "Medico", "Sector 1", 3, 1, 0, 1, RUT])
    ws.add_table(Table(displayName="Tabla1", ref="A7:C36"))
    wb.create_sheet("Hoja2")
    wb.create_sheet("Hoja3")
    ruta = _TMP / nombre
    wb.save(ruta)
    return ruta


def test_el_script_arranca_en_el_python_del_proyecto():
    # `int | None` sin `from __future__ import annotations` revienta al IMPORTAR en 3.9.
    r = subprocess.run([sys.executable, str(REPO / "tools" / "limpiar_refs.py"), "no_existe.xlsx"],
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    assert r.returncode == 0 and "no existe" in r.stdout, (r.returncode, r.stderr[-300:])


def test_deja_banner_y_encabezado_de_dos_pisos_y_nada_mas():
    ruta = _crudo("cupos.xlsx")
    estado, antes, despues, _ = L.recortar(ruta, aplicar=True)
    assert estado == "recortado", estado
    assert (antes, despues) == (36, 6), (antes, despues)
    wb = openpyxl.load_workbook(ruta)
    ws = wb["Hoja1"]
    assert ws["A1"].value == "Comuna" and ws["B2"].value == "[CESFAM] Ejemplo"   # banner
    assert ws["E6"].value == "Total"                    # 2o piso del encabezado
    assert ws.max_row == 6, ws.max_row                  # ni una fila de datos
    assert {str(m) for m in ws.merged_cells.ranges} == {"A5:A6", "E5:G5"}
    assert wb.sheetnames == ["Hoja1", "Hoja2", "Hoja3"]  # la hoja unica + 2 vacias
    with zipfile.ZipFile(ruta) as z:
        assert not [n for n in z.namelist() if "table" in n.lower()], z.namelist()
    assert not escanear(ruta)[0]
    assert L.recortar(ruta, aplicar=False)[0] == "ya limpio"


def test_un_banner_con_rut_no_se_escribe():
    ruta = _crudo("con_rut.xlsx", banner_extra=["Generado para:", RUT_CAZABLE])
    antes = ruta.read_bytes()
    estado, *_r, avisos = L.recortar(ruta, aplicar=True)
    assert estado == "NO ESCRITO", estado
    assert any("RUT" in a for a in avisos), avisos
    assert RUT_CAZABLE not in " ".join(avisos)                  # el aviso cuenta, no transcribe
    assert ruta.read_bytes() == antes                   # el archivo queda como estaba
    assert not list(_TMP.glob("*.revisando.xlsx"))


def test_las_referencias_versionadas_estan_limpias():
    tracked = subprocess.run(["git", "ls-files", "refs_tablas/*.xlsx"], cwd=REPO,
                             capture_output=True, text=True).stdout.split()
    revisadas = 0
    for rel in tracked:
        p = REPO / rel
        if not L._es_referencia_export(p):
            continue
        estado, *_r, avisos = L.recortar(p, aplicar=False)
        assert estado == "ya limpio", f"{rel}: {estado} {avisos}"
        assert not escanear(p)[0], f"{rel}: el escaner encuentra RUT/email/telefono"
        revisadas += 1
    assert revisadas >= 10, revisadas


def _literales_de_actividad(ruta):
    """Cada literal que un modulo busca en la ACTIVIDAD: los `_all(A, ...)` /
    `_any(A, [...])` (A = la serie de actividades), leidos del codigo por AST para que
    un filtro nuevo quede cubierto sin tocar este test."""
    import ast
    out = set()
    for nodo in ast.walk(ast.parse(Path(ruta).read_text(encoding="utf-8"))):
        if (isinstance(nodo, ast.Call) and getattr(nodo.func, "id", "") in ("_all", "_any")
                and nodo.args and getattr(nodo.args[0], "id", "") == "A"):
            for arg in nodo.args[1:]:
                for c in ast.walk(arg):
                    if isinstance(c, ast.Constant) and isinstance(c.value, str):
                        out.add(c.value)
    return out


def test_cada_actividad_que_se_busca_existe_en_el_maestro():
    """El bug del A32-F2 (1.9.10) y de 3 de las 7 actividades del «Activo 12m» (ronda
    11): un literal que no calza con NINGUNA actividad real da 0 para siempre, con cara
    de «ese mes no hubo». El Maestro de Actividades versionado es el catalogo real."""
    from programas.rem_utils import cargar_maestro, norm
    from programas.poblacion import ACTIVIDADES_SM_7
    from modulos.rem_sm_actividades import ADA_TRIBUTAN
    from modulos.rem_sm_trabajo_perdido import EXCLUIR_SMISH
    acts = set(cargar_maestro(REPO / "catalogos" / "maestro_slim.csv.gz")["ACT_n"])
    buscados = {
        "poblacion.ACTIVIDADES_SM_7": set(ACTIVIDADES_SM_7),
        "rem_sm_actividades.ADA_TRIBUTAN": set(ADA_TRIBUTAN),
        "rem_sm_trabajo_perdido.EXCLUIR_SMISH": {p for p, _ in EXCLUIR_SMISH},
        "rem_sm_actividades (_all/_any)": _literales_de_actividad(REPO / "modulos" / "rem_sm_actividades.py"),
        "rem_a23_respiratorio (_all/_any)": _literales_de_actividad(REPO / "modulos" / "rem_a23_respiratorio.py"),
    }
    assert len(buscados["rem_a23_respiratorio (_all/_any)"]) > 15, buscados
    muertos = {f"{donde}: {lit!r}" for donde, lits in buscados.items() for lit in lits
               if not any(norm(lit) in a for a in acts)}
    assert not muertos, "no calzan con NINGUNA actividad del Maestro:\n" + "\n".join(sorted(muertos))


def _main():
    pruebas = [v for k, v in sorted(globals().items())
               if k.startswith("test_") and callable(v)]
    fallos = 0
    for fn in pruebas:
        try:
            fn(); print(f"PASS  {fn.__name__}")
        except AssertionError as e:
            fallos += 1; print(f"FAIL  {fn.__name__} -> {e or 'assert'}")
        except Exception as e:  # noqa: BLE001
            fallos += 1; print(f"ERROR {fn.__name__} -> {type(e).__name__}: {e}")
    print("-" * 50)
    print(f"{len(pruebas) - fallos}/{len(pruebas)} OK" + (f" ({fallos} con problemas)" if fallos else ""))
    return 1 if fallos else 0


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    sys.exit(_main())
