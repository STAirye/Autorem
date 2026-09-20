#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# ==========================================================================
# This code was generated with the assistance of Claude Sonnet 5 (Anthropic).
# The human author reviewed, modified, and integrated the code.
#
# Author: Simón Tobar — CESFAM Dr. Luis Ferrada Urzúa (APS, SSMC)
# Copyright (C) 2026 Simón Tobar
# SPDX-License-Identifier: GPL-3.0-or-later
# ==========================================================================
"""
Pruebas del REGISTRO de pantallas de la GUI 2.0 (gui/registro.py + gui/app.py,
docs/GUI_2.0_plan.md §11).

Anti-olvido, mismo espíritu que tests/test_cobertura.py: descubre por
introspección las páginas de gui/paginas/*.py y falla si una PANTALLA queda
mal declarada — que es justo donde una migración de este tamaño pierde cosas
en silencio. NO prueba que la ventana se vea bien (eso es a ojo, §13); prueba
que el CONTRATO esté completo.

Correr:
    python tests/test_gui_registro.py    # runner propio, imprime PASS/FAIL
    pytest tests/                         # si tienes pytest instalado
"""

import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
import _aislar_cache   # noqa: E402,F401  (PRIMERO: nunca tocar el ~/.autorem real)

from gui.registro import cargar_registro, ORDEN_PAGINAS, ORDEN_PROGRAMAS   # noqa: E402

_CLAVES_OBLIGATORIAS = ("id", "programa", "titulo", "estado", "inputs", "mes",
                        "carpeta_salida", "correr", "resumen")


def _registro():
    r = cargar_registro()
    assert r, "gui/paginas/ no expuso ninguna PANTALLA -- ¿import roto?"
    return r


def test_claves_obligatorias():
    for pantalla in _registro():
        faltantes = [k for k in _CLAVES_OBLIGATORIAS if k not in pantalla]
        assert not faltantes, f"{pantalla.get('id', '?')}: faltan claves {faltantes}"


def test_ids_unicos():
    ids = [p["id"] for p in _registro()]
    dup = {i for i in ids if ids.count(i) > 1}
    assert not dup, f"id(s) de PANTALLA repetidos: {dup}"


def test_programa_declarado_en_el_orden_del_sidebar():
    # `registro._orden_programa` deja pasar un programa desconocido (fail
    # SOFT: cae al final, no revienta el sidebar) -- este test es el fail
    # LOUD que le falta: un programa nuevo hay que declararlo a mano.
    for pantalla in _registro():
        assert pantalla["programa"] in ORDEN_PROGRAMAS, (
            f"{pantalla['id']}: programa {pantalla['programa']!r} no está en "
            f"registro.ORDEN_PROGRAMAS -- agrégalo ahí (si no, el sidebar lo "
            f"manda al final en silencio)")


def test_pagina_declarada_en_el_orden_del_sidebar():
    """Hermano del de arriba, para el orden DENTRO del grupo. Sin ORDEN_PAGINAS la
    posicion la decidia `pkgutil`, o sea el NOMBRE DEL ARCHIVO: renombrar sm.py a
    actividades.py reordenaba el sidebar en silencio. `_orden_pagina` deja pasar una
    pagina no declarada (fail SOFT: al final del grupo) -- este es el fail LOUD."""
    for pantalla in _registro():
        assert pantalla["id"] in ORDEN_PAGINAS, (
            f"{pantalla['id']}: no está en registro.ORDEN_PAGINAS -- agrégalo ahí "
            f"(si no, el sidebar lo manda al final de su grupo en silencio)")


def test_el_orden_de_paginas_no_depende_del_nombre_del_archivo():
    """El de arriba asegura que la pagina ESTE declarada; este, que la declaracion se
    APLIQUE. Hace falta un registro sintetico: hoy el orden alfabetico de los .py
    coincide por casualidad con el deseado ("a05" < "sm"), asi que el registro real no
    puede distinguir "ordena por ORDEN_PAGINAS" de "ordena por nombre de archivo" --
    que es justo la confusion que dejo el bug. Aca `_paginas_encontradas` devuelve las
    paginas AL REVES y el orden declarado tiene que ganar igual."""
    from gui import registro as reg

    falsas = [{"id": pid, "programa": "Salud Mental"}
              for pid in ("sm_actividades", "a05")]     # al reves a proposito
    previo = reg._paginas_encontradas
    reg._paginas_encontradas = lambda: list(falsas)
    try:
        ids = [p["id"] for p in reg.cargar_registro()]
    finally:
        reg._paginas_encontradas = previo
    assert ids == ["a05", "sm_actividades"], (
        f"el sidebar quedo en {ids}: el orden sale del orden en que se ENCONTRARON "
        f"los archivos, no de registro.ORDEN_PAGINAS")


def test_una_sola_ancla_de_salida_por_pantalla():
    """`ancla_salida` fija la carpeta de salida por defecto (gui/app.py
    `_resolver_ctx`). Dos anclas en una pantalla = el default lo decide otra vez el
    orden en que estan ESCRITOS los inputs, que es justo lo que la marca elimina."""
    for pantalla in _registro():
        anclas = [i["key"] for i in pantalla.get("inputs", []) if i.get("ancla_salida")]
        assert len(anclas) <= 1, f"{pantalla['id']}: más de un ancla_salida: {anclas}"


def test_correr_y_resumen_invocables():
    for pantalla in _registro():
        assert callable(pantalla["correr"]), f"{pantalla['id']}: 'correr' no es invocable"
        assert callable(pantalla["resumen"]), f"{pantalla['id']}: 'resumen' no es invocable"


def test_preparar_y_al_completar_invocables_si_estan():
    for pantalla in _registro():
        if pantalla.get("preparar") is not None:
            assert callable(pantalla["preparar"]), f"{pantalla['id']}: 'preparar' no es invocable"
        if pantalla.get("al_completar") is not None:
            assert callable(pantalla["al_completar"]), (
                f"{pantalla['id']}: 'al_completar' no es invocable")


def test_inputs_key_no_duplicada():
    for pantalla in _registro():
        keys = [inp["key"] for inp in pantalla.get("inputs", [])]
        dup = {k for k in keys if keys.count(k) > 1}
        assert not dup, f"{pantalla['id']}: input key(s) repetida(s) {dup}"


def test_extras_despues_de_apunta_a_un_input_real():
    for pantalla in _registro():
        keys_inputs = {inp["key"] for inp in pantalla.get("inputs", [])}
        for extra in pantalla.get("extras", []):
            despues_de = extra.get("despues_de")
            if despues_de is not None:
                assert despues_de in keys_inputs, (
                    f"{pantalla['id']}: extras[].despues_de={despues_de!r} no es la key "
                    f"de ningún input de esta página (keys: {sorted(keys_inputs)})")


class _MessageboxFalso:
    """Captura los avisos en vez de abrir una ventana: `valida_mes` es pura salvo
    por el messagebox, asi que se puede probar sin display."""

    def __init__(self):
        self.avisos = []

    def showwarning(self, titulo, mensaje):
        self.avisos.append((titulo, mensaje))


def test_valida_mes_rechaza_el_rango_no_solo_lo_no_numerico():
    """El `ttk.Spinbox` acota SOLO sus flechas: un mes 13 TECLEADO llegaba al worker
    y reventaba en `_rango_mes` con 'month must be in 1..12', que `manejar_error`
    despacha como "Error inesperado ... pasaselo a Simon" -- un typo presentado como
    bug del programa, y despues de cargar los archivos. Un año de 2 digitos era peor:
    no reventaba, salia 'no hay filas de 07/0026' culpando al export.

    `_tab_a05` y `_tab_beta` validaban el rango en autorem.py; el port lo perdio para
    todas las paginas menos A05. Ahora es una sola funcion compartida."""
    from gui.runner import valida_mes
    from gui.widgets import ANIO_MAX, ANIO_MIN

    for malo in (None, (2026, 13), (2026, 0), (2026, -1), (26, 7), (202, 7),
                 (ANIO_MIN - 1, 6), (ANIO_MAX + 1, 6)):
        mb = _MessageboxFalso()
        assert valida_mes(malo, mb) is None, f"{malo} deberia rechazarse"
        assert mb.avisos, f"{malo} se rechazo en SILENCIO (sin avisar al usuario)"

    for bueno in ((2026, 1), (2026, 12), (2026, 7), (ANIO_MIN, 1), (ANIO_MAX, 12)):
        mb = _MessageboxFalso()
        assert valida_mes(bueno, mb) == bueno, f"{bueno} es valido y se rechazo"
        assert not mb.avisos, f"{bueno} es valido y aun asi aviso: {mb.avisos}"


def test_el_spinbox_y_la_guarda_de_anio_no_pueden_divergir():
    """Los `from_`/`to` del Spinbox y el rango que valida `runner.valida_mes` salen de
    las MISMAS constantes. Si alguien mueve una sola, el mensaje de error miente."""
    import inspect
    from gui import widgets
    fuente = inspect.getsource(widgets.selector_mes)
    assert "from_=ANIO_MIN, to=ANIO_MAX" in fuente, (
        "selector_mes volvio a hardcodear los años del Spinbox: tiene que usar "
        "widgets.ANIO_MIN/ANIO_MAX, que es lo que valida runner.valida_mes")


def test_motivo_fuente_sigue_las_mismas_ramas_que_manejar_error():
    """El BannerFuente (al elegir el archivo) y el messagebox (al procesar) describen
    el MISMO archivo, asi que no pueden contradecirse: los dos salen del arbol de
    `isinstance` de `manejar_error`. Un `.xls` disfrazado de `.xlsx` no es un problema
    de firmas y no puede decir que lo sea."""
    import zipfile
    from gui.runner import es_error_formato, motivo_fuente

    disfrazado = zipfile.BadZipFile("File is not a zip file")
    assert es_error_formato(disfrazado), "BadZipFile deberia ser problema de formato"
    assert "xlsx" in motivo_fuente(disfrazado)
    assert "Excel" in motivo_fuente(PermissionError(13, "in use"))
    assert "librería" in motivo_fuente(ImportError("Falta 'openpyxl'"))
    # Cualquier otra cosa: se nombra el tipo, no se inventa una causa.
    assert "ValueError" in motivo_fuente(ValueError("cualquiera"))


def test_el_desglose_por_instrumento_del_a03_llega_hasta_el_resumen():
    """Los 3 slots A03·D.3 fijan el instrumento A MANO (sin autodeteccion), asi que un
    export cargado en el slot equivocado solo se delata en el desglose: un total pelado
    ("60 aplicaciones") tapa igual un 20/20/20 que un 60/0/0. La vieja pestana
    standalone lo mostraba y el port lo perdio.

    Corre `correr` de verdad en modo solo-cuestionarios (con `procesar_unificado` y los
    estamentos stubbeados) para probar TODO el camino: que el dato se guarde en `res` y
    que `resumen` lo imprima. Probar solo `_por_instrumento` con un dict a mano dejaba
    pasar que nadie lo estuviera poblando."""
    import modulos.rem_a03_d3_instrumentos as screening
    import programas.estamentos as estam
    from gui.paginas import sm

    import tempfile
    por_inst = {"PSC": 20, "PSC-Y": 18, "GHQ-12": 22}
    previos = (screening.procesar_unificado, estam.tabla_efectiva)

    def unificado_falso(por_instrumento, salida, **kw):
        Path(salida).write_bytes(b"d3")   # escribe como el real (lo renombra escribir_atomico)
        return {"salida": str(salida), "total": 60, "por_instrumento": por_inst, "tabla": None,
                "avisos": [("Tabla D.3", "SUBCONTADO", "2 fuera de rango", "revisar")]}
    screening.procesar_unificado = unificado_falso
    estam.tabla_efectiva = lambda *_a, **_kw: None
    try:
        # Carpeta TEMPORAL: la corrida escribe, y nunca dentro del repo.
        ctx = {"mes": (2026, 8), "carpeta": Path(tempfile.mkdtemp(prefix="autorem_a03_")),
               "solo_a03": True,
               "a03": {"incluir": True, "est_ruta": "",
                       "instrumentos": {k: f"{k}.xlsx" for k in por_inst}}}
        res = sm.correr(ctx, log=lambda *_a: None)
    finally:
        screening.procesar_unificado, estam.tabla_efectiva = previos

    assert res["por_inst_a03"] == por_inst, "el desglose no llego a `res`"
    texto = sm.resumen(res)
    # Ronda 10: los avisos del A03 tambien tienen que llegar (antes `correr` los botaba).
    for trozo in ("60 aplicaciones", "PSC: 20", "PSC-Y: 18", "GHQ-12: 22", "2 fuera de rango"):
        assert trozo in texto, f"falta {trozo!r} en el resumen: {texto!r}"
    # Sin dato no se inventa un parentesis vacio.
    assert sm._por_instrumento({"por_inst_a03": None}) == ""
    assert sm._por_instrumento({}) == ""


def test_sm_no_pisa_salidas_y_el_resumen_dice_lo_que_no_se_genero():
    """Corrida COMPLETA de SM (Actividades + A03) en una carpeta que ya tiene el A03·D.3
    de una corrida anterior del mismo mes, y el A03 de esta falla. Antes: el A03 viejo
    quedaba junto a la salida nueva con su mismo nombre, el fallo solo se veia en el
    log, y el 'Listo' no decia nada -> el usuario abria el archivo viejo y copiaba al
    REM una D.3 de otros cuestionarios. Ahora (1) toda la corrida sale como `(1)`, (2)
    el viejo queda intacto, y (3) el resumen dice que la D.3 NO se genero."""
    import tempfile
    import pandas as pd
    import modulos.rem_sm_actividades as smact
    import modulos.rem_sm_trabajo_perdido as tpmod
    import programas.rem_utils as ru
    import modulos.rem_a03_d3_instrumentos as screening
    import programas.estamentos as estam
    from gui.paginas import sm

    carpeta = Path(tempfile.mkdtemp(prefix="autorem_sm_"))
    viejo = carpeta / "REM_A03_D3_2026_08.xlsx"
    viejo.write_bytes(b"D.3 de otra corrida")

    E = pd.DataFrame({"x": [1]})
    E.attrs["tablas"] = {"SM_Resumen": pd.DataFrame({"Casilla": ["A04"], "Total mes": [1]})}

    def a03_falla(*_a, **_k):
        raise ValueError("instrumento raro")

    atomico, atomicos = ru.escribir_atomico, []
    parches = [(ru, "escribir_atomico",
                lambda s, fn: (atomicos.append(Path(s).name), atomico(s, fn))[1]),
               (smact, "procesar", lambda *_a, **_k: E),
               (smact, "escribir", lambda _E, salida: Path(salida).write_bytes(b"nuevo")),
               (tpmod, "procesar", lambda *_a, **_k: pd.DataFrame()),
               (ru, "cargar_maestro", lambda *_a, **_k: None),   # el Maestro de mentira
               (tpmod, "escribir", lambda _E, salida: Path(salida).write_bytes(b"nuevo")),
               (screening, "procesar_unificado", a03_falla),
               (estam, "tabla_efectiva", lambda *_a, **_k: None)]
    previos = [(m, n, getattr(m, n)) for m, n, _ in parches]
    for m, n, v in parches:
        setattr(m, n, v)
    try:
        ctx = {"mes": (2026, 8), "carpeta": carpeta, "solo_a03": False,
               "ada": [Path("ada.xlsx")], "grupal": [], "inscritos": None,
               "multiprofesional": None, "maestro": Path("maestro.csv"),
               "d": None, "tabla_dot": None,
               "a03": {"incluir": True, "est_ruta": "", "instrumentos": {"PSC": "psc.xlsx"}}}
        res = sm.correr(ctx, log=lambda *_a: None)
    finally:
        for m, n, v in previos:
            setattr(m, n, v)

    assert atomicos == ["REM_SM_actividades_2026_08 (1).xlsx",
                        "REM_SM_trabajo_perdido_2026_08 (1).xlsx",
                        "REM_A03_D3_2026_08 (1).xlsx"], (
        f"alguna salida no se escribio via temporal (escribir_atomico): {atomicos}")
    assert res["salida"].name == "REM_SM_actividades_2026_08 (1).xlsx", (
        f"la corrida no tomo el mismo `(1)` que el A03 que ya existia: {res['salida'].name}")
    assert viejo.read_bytes() == b"D.3 de otra corrida", "piso el A03 de la corrida anterior"
    texto = sm.resumen(res)
    assert "NO se generó" in texto and "instrumento raro" in texto, (
        f"el resumen calla que la D.3 no salio: {texto!r}")


def test_a23_y_poblacion_tampoco_pisan_salidas():
    """Mismo arreglo que SM, en las otras dos paginas que escriben: nada se pisa, y el
    Rescate que falla se dice en el resumen de Poblacion (no solo en el log)."""
    import tempfile
    import pandas as pd
    import modulos.rem_a23_respiratorio as a23m
    import programas.poblacion as pob
    import modulos.rem_sp_p6_poblacion as p6
    import modulos.rem_sm_rescate_inasistentes as resc
    import programas.rem_utils as ru
    from gui.paginas import a23, poblacion

    carpeta = Path(tempfile.mkdtemp(prefix="autorem_pag_"))
    (carpeta / "REM_A23_2026_08_procesado.xlsx").write_bytes(b"viejo")
    (carpeta / "REM_SM_Rescate_2026_08_BETA.xlsx").write_bytes(b"viejo")
    P = pd.DataFrame({"¿Ingresado?": ["SI"]})

    def rescate_falla(*_a, **_k):
        raise ValueError("brecha rara")

    escribe = lambda *a: Path(a[-1]).write_bytes(b"nuevo")   # noqa: E731
    atomico, atomicos = ru.escribir_atomico, []
    parches = [(ru, "escribir_atomico",
                lambda s, fn: (atomicos.append(Path(s).name), atomico(s, fn))[1]),
               (a23m, "procesar", lambda *_a, **_k: pd.DataFrame({"x": [1]})),
               (a23m, "escribir", escribe),
               (pob, "cargar_inscritos", lambda *_a, **_k: None),
               (pob, "cargar_formulario_sm", lambda *_a, **_k: None),
               (ru, "cargar_atenciones", lambda *_a, **_k: None),
               (pob, "construir_poblacion", lambda *_a, **_k: P),
               (p6, "construir_p6", lambda *_a, **_k: {"revisar_administrativo": [],
                                                      "revisar_clinico": []}),
               (p6, "escribir", escribe),
               (resc, "procesar", rescate_falla)]
    previos = [(m, n, getattr(m, n)) for m, n, _ in parches]
    for m, n, v in parches:
        setattr(m, n, v)
    try:
        ra = a23.correr({"mes": (2026, 8), "carpeta": carpeta, "atenciones": [],
                         "otros_cronicos": [], "estratificacion": [], "nsp": []},
                        log=lambda *_a: None)
        rp = poblacion.correr({"mes": (2026, 8), "carpeta": carpeta,
                               "inscritos": Path("insc.xlsx"), "formularios": [], "ada": []},
                              log=lambda *_a: None)
    finally:
        for m, n, v in previos:
            setattr(m, n, v)

    assert ra["salida"].name == "REM_A23_2026_08_procesado (1).xlsx", ra["salida"].name
    # El Rescate revienta DENTRO de procesar, antes de escribir: solo A23 y P6 pasan.
    assert atomicos == ["REM_A23_2026_08_procesado (1).xlsx",
                        "REM_SP_P6_2026_08_BETA (1).xlsx"], (
        f"alguna salida no se escribio via temporal (escribir_atomico): {atomicos}")
    # El Rescate viejo ocupa el nombre -> el P6 de ESTA corrida tambien va como (1).
    assert rp["salida"].name == "REM_SP_P6_2026_08_BETA (1).xlsx", rp["salida"].name
    assert (carpeta / "REM_SM_Rescate_2026_08_BETA.xlsx").read_bytes() == b"viejo"
    assert "NO se generó" in poblacion.resumen(rp), "el resumen calla que el Rescate fallo"


def test_un_opcional_invalido_pregunta_si_seguir_sin_el():
    """Ronda 11 (decision del autor): un archivo OPCIONAL que no sirve no se calla (TRANS
    y el Multiprofesional quedaban en el log) ni tumba la corrida: la app pregunta
    «¿continuar sin el?»; si si, re-corre sin el y la LEEME dice que se omitio."""
    import tempfile
    import pandas as pd
    import modulos.rem_a23_respiratorio as a23m
    from programas.rem_utils import ArchivoInvalido, OpcionalInvalido
    from gui import runner
    from gui.paginas import a23

    class _Msg:
        def __init__(self, resp): self.resp, self.preguntas = resp, []
        def askyesno(self, titulo, texto): self.preguntas.append(texto); return self.resp

    inputs = a23.PANTALLA["inputs"]
    # La Estratificacion es un input de UN archivo (`multi: False` desde la ronda 12):
    # `sin_opcional` lo deja en None, no en [].
    ctx = {"mes": (2026, 8), "carpeta": Path(tempfile.mkdtemp(prefix="autorem_opc_")),
           "atenciones": [Path("a.xlsx")], "otros_cronicos": [Path("o.xlsx")],
           "estratificacion": Path("e.xlsx"), "nsp": []}
    err = OpcionalInvalido("estrat", ArchivoInvalido("sin_columnas", "no encuentro RUT"))
    q = lambda *_a: None   # noqa: E731

    si = _Msg(True)
    nuevo = runner.sin_opcional(err, ctx, inputs, si, q)
    assert nuevo["estratificacion"] is None and ctx["estratificacion"], "no quito el archivo (o toco el ctx original)"
    assert nuevo["descartados"][0][0].startswith("Estratificación") and "RUT" in si.preguntas[0]
    assert runner.sin_opcional(err, ctx, inputs, _Msg(False), q) is False      # dijo que no
    assert runner.sin_opcional(ArchivoInvalido("x", "y"), ctx, inputs, si, q) is None   # no es opcional
    oblig = OpcionalInvalido("otros_cronicos", ArchivoInvalido("x", "y"))
    assert runner.sin_opcional(oblig, ctx, inputs, si, q) is None             # obligatorio: no se ofrece

    fer = pd.DataFrame({"x": [1]})
    fer.attrs["avisos"] = []
    previos = (a23m.procesar, a23m.escribir)
    a23m.procesar = lambda *_a, **_k: fer
    a23m.escribir = lambda _f, p: Path(p).write_bytes(b"x")
    try:
        res = a23.correr(nuevo, log=q)
    finally:
        a23m.procesar, a23m.escribir = previos
    assert any(a[1] == "OMITIDO" and "Estratificación" in a[0] for a in fer.attrs["avisos"]), fer.attrs
    assert "OMITIDO" in a23.resumen(dict(res, fer=fer)) or "Estratificación" in a23.resumen(dict(res, fer=fer))


def test_un_click_encolado_no_abre_una_segunda_ventana_de_dotacion():
    """`ctk.CTkToplevel(root)` hace un `update()` completo en su constructor (Windows),
    asi que un click que quedo encolado mientras se cargaba el ADA se despacha DENTRO
    del dialogo de dotacion de `preparar`. Sin candado, un «Revisar»/«Precargar
    dotación…» se abria anidado con su propia copia de la tabla, y el 'Aplicar' de
    afuera pisaba lo guardado adentro (y la corrida usaba la tabla vieja).

    El dialogo se reemplaza por uno falso que hace lo mismo que el constructor de
    CTkToplevel: despachar esos clicks. La premisa (que CTkToplevel de verdad los
    despacha) la amarra `test_gui_construccion`."""
    import pandas as pd
    import programas.rem_utils as ru
    from programas import dotacion
    from gui import dialogos

    lecturas, logs, anidadas = [], [], {}
    ada = pd.DataFrame({"ACT_n": ["X"]})
    fila = pd.DataFrame({"funcionario": ["ANA SOTO"], "estamento": ["PSICOLOGO"],
                         "n_atenciones": [3]})

    def dialogo_falso(root, tabla, modulo, filas, **_kw):
        if anidadas:
            return
        anidadas["en_curso"] = True
        anidadas["precargar"] = dialogos.dotacion_ada(
            None, "sm", ["ada.xlsx"], (2026, 8), logs.append, None, todos=True)
        dialogos.revisar_dotacion(None, log=logs.append)

    parches = [
        (ru, "cargar_atenciones", lambda *_a, **_k: (lecturas.append("ada"), ada)[1]),
        (ru, "filtrar_mes", lambda d, *_a, **_k: d),
        (dotacion, "cargar", lambda *_a, **_k: {"funcionarios": {}, "omitidos": {}}),
        (dotacion, "evidencia", lambda *_a, **_k: fila),
        (dotacion, "nuevos", lambda ev, *_a: ev),
        (dialogos, "dialogo_dotacion", dialogo_falso),
        (dialogos, "_revisar_dotacion", lambda *_a, **_k: lecturas.append("revisar")),
    ]
    previos = [(m, n, getattr(m, n)) for m, n, _ in parches]
    for m, n, v in parches:
        setattr(m, n, v)
    try:
        d, tabla = dialogos.dotacion_ada(None, "sm", ["ada.xlsx"], (2026, 8),
                                         logs.append, None)
    finally:
        for m, n, v in previos:
            setattr(m, n, v)

    assert anidadas.get("en_curso"), "el dialogo falso no se abrio: el test no prueba nada"
    assert lecturas == ["ada"], (
        f"un click encolado entro a otra ventana de dotacion anidada: {lecturas}")
    assert anidadas["precargar"] == (None, None), "el Precargar anidado no aborto"
    # Los DOS abortos (Precargar y Revisar) se dicen: Revisar lo callaba.
    assert sum("ya hay una ventana" in m for m in logs) == 2, (
        f"un aborto no se dijo en el log: {logs}")
    assert d is ada and tabla is not None, "la ventana de afuera no termino normal"
    assert dialogos._DOTACION_ABIERTA[0] is False, "el candado quedo tomado"


class _MbAvisos:
    def __init__(self):
        self.vistos = []

    def showwarning(self, titulo, mensaje):
        self.vistos.append((titulo, mensaje))


def test_los_avisos_de_cache_se_muestran_al_cerrar_los_dialogos_de_dotacion():
    """Un caché dañado o que no se pudo guardar tiene que VERSE (no solo el log, que en
    el exe --windowed no ve nadie): `avisar_cache` junta todos en UN dialogo, y los
    dialogos de dotacion (Precargar/Procesar y Revisar) lo llaman al cerrarse."""
    import tkinter.messagebox as tkmb
    import programas.rem_utils as ru
    from gui import dialogos, runner

    ru.tomar_avisos_cache()
    mb = _MbAvisos()
    runner.avisar_cache(mb)
    assert mb.vistos == [], "sin avisos no se muestra nada"

    def empuja(*_a, **_k):
        ru._avisar_cache(lambda *_: None, "No pude guardar la tabla de dotación")
        return None, None
    previos = (dialogos._dotacion_ada, dialogos._revisar_dotacion, tkmb.showwarning)
    revisar = _MbAvisos()
    dialogos._dotacion_ada = empuja
    dialogos._revisar_dotacion = empuja
    tkmb.showwarning = revisar.showwarning
    try:
        dialogos.dotacion_ada(None, "sm", [], (2026, 8), lambda *_: None, mb)
        dialogos.revisar_dotacion(None)
    finally:
        dialogos._dotacion_ada, dialogos._revisar_dotacion, tkmb.showwarning = previos
    assert len(mb.vistos) == 1 and "No pude guardar" in mb.vistos[0][1], mb.vistos
    assert len(revisar.vistos) == 1, "Revisar dotación no mostro el aviso de caché"
    assert ru.tomar_avisos_cache() == [], "quedaron avisos sin mostrar"


def test_sm_valida_la_ruta_de_cupos_antes_del_worker():
    """'Utilización de Cupos' es opcional, pero una ruta mal tecleada tiene que
    pararse en `preparar` (hilo GUI, con dialogo). Antes recien reventaba al final
    de la corrida, con SM y TP ya escritos y la D.3 perdida en `fallo_a03`."""
    import tkinter.messagebox as tkmb
    from gui.paginas import sm

    vistos = []
    previo = tkmb.showerror
    tkmb.showerror = lambda t, m: vistos.append((t, m))
    try:
        ctx = {"ada": [], "grupal": [], "mes": (2026, 8),
               "a03": {"incluir": True, "instrumentos": {"PSC": "psc.xlsx"},
                       "est_ruta": "C:/no/existe/cupos.xlsx"}}
        assert sm.preparar(ctx, pagina=None) is None, "siguio con una ruta de Cupos inexistente"
        assert vistos and "cupos.xlsx" in vistos[0][1], vistos
        # Sin ruta (opcional) sigue normal.
        ctx["a03"]["est_ruta"] = ""
        assert sm.preparar(ctx, pagina=None) is ctx
    finally:
        tkmb.showerror = previo


def test_el_preview_de_cruce_lee_el_encabezado_con_el_criterio_del_loader():
    """`sm._header_rapido` (preview ADA<->Grupal) tiene que elegir la MISMA fila que
    `leer_xlsx`. Era una copia a mano del criterio: si el loader cambiaba, el preview
    leia otro encabezado. Se prueba el cableado: con el criterio compartido
    parchado, el preview lo sigue."""
    import tempfile
    import openpyxl
    import programas.rem_utils as ru
    from gui.paginas import sm

    ruta = Path(tempfile.mkdtemp(prefix="autorem_hdr_")) / "ada.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Servicio de Salud"])
    ws.append(["A", "B", "C", "D", "E"])
    ws.append([1, 2, 3, 4, 5])
    wb.save(ruta)

    assert sm._header_rapido(ruta) == ru.leer_xlsx(ruta)[0] == ["A", "B", "C", "D", "E"]
    previo = ru.indice_encabezado
    ru.indice_encabezado = lambda filas, *a, **k: 2
    try:
        assert sm._header_rapido(ruta) == [1, 2, 3, 4, 5], (
            "el preview no usa rem_utils.indice_encabezado (criterio copiado a mano)")
    finally:
        ru.indice_encabezado = previo


def test_el_periodo_del_a05_es_el_selector_mes_compartido():
    """La caja Periodo del A05 tenia su propia copia de los dos Spinbox y del
    parseo, fuera del test que amarra los años a `valida_mes`. Ahora usa
    `widgets.selector_mes`."""
    import inspect
    from gui.paginas import a05
    fuente = inspect.getsource(a05.bloque_periodo)
    assert "widgets.selector_mes(" in fuente, "el A05 no usa selector_mes"
    assert "ttk.Spinbox(" not in fuente, "el A05 volvio a armar sus Spinbox a mano"
    assert "int(" not in inspect.getsource(a05.preparar), "el A05 volvio a parsear el mes a mano"


def test_aplicar_dotacion_solo_manda_lo_que_cambio():
    """«Precargar» y «Revisar» muestran gente YA clasificada. Si 'Aplicar' mandaba
    todos los ticks, reescribia la foto de cuando se abrio la ventana sobre lo que
    otra ventana hubiera guardado entremedio. Un nombre NUEVO sin tick si se manda:
    quedar interno es una decision."""
    from programas import dotacion
    from gui import dialogos
    tabla = {"funcionarios": {"ANA SOTO": dotacion.INTERNO, "BETO RUIZ": dotacion.EXTERNO},
             "omitidos": {}}
    ticks = {"ANA SOTO": False, "BETO RUIZ": False, "CARLA PAZ": False}
    assert dialogos.decisiones_cambiadas(ticks, tabla) == {"BETO RUIZ": False, "CARLA PAZ": False}


def test_dotacion_dice_que_nada_tributa_en_vez_de_todo_en_orden():
    """Ronda 9: con un ADA del mes en el que NADA tributa al REM, el camino de
    Procesar (todos=False) logueaba «sin funcionarios nuevos que clasificar», que se
    lee como "ya estan todos clasificados". Tiene que decir que nada tributa."""
    import openpyxl
    import tempfile
    import modulos.rem_sm_actividades as smact
    from gui import dialogos
    p = Path(tempfile.mkdtemp(prefix="autorem_dot_")) / "ada.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["NUMERO TIPO IDENTIFICACION", "ATEN ID", "FECHA ATENCION", "ACTIVIDADES",
               "DIAGNOSTICOS", "INSTRUMENTO", "TIPO ATENCION", "FUNCIONARIO"])
    ws.append(["Z", "Z1", "15/07/2026", "Curacion simple", "", "Enfermero(a)", "Consulta",
               "ANA SOTO"])
    wb.save(p)

    class _MB:
        def __getattr__(self, _):
            return lambda *a, **k: None
    logs = []
    d, _tabla = dialogos._dotacion_ada(None, "sm", [p], (2026, 7), logs.append, _MB(),
                                       mask=smact.mask_tributa_ada)
    assert d is not None
    assert any("NINGUNA atención" in m for m in logs), logs
    assert not any("sin funcionarios nuevos" in m for m in logs), logs

    # 2a pasada: SI tributa, pero sin FUNCIONARIO -> tampoco es "todo en orden"
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["NUMERO TIPO IDENTIFICACION", "ATEN ID", "FECHA ATENCION", "ACTIVIDADES",
               "DIAGNOSTICOS", "INSTRUMENTO", "TIPO ATENCION", "FUNCIONARIO"])
    ws.append(["A", "A1", "15/07/2026", "Consulta De Salud Mental", "", "Médico", "Consulta", ""])
    wb.save(p)
    logs.clear()
    dialogos._dotacion_ada(None, "sm", [p], (2026, 7), logs.append, _MB(), mask=smact.mask_tributa_ada)
    assert any("NO traen FUNCIONARIO" in m for m in logs), logs
    assert not any("sin funcionarios nuevos" in m for m in logs), logs


def test_construye_todas_las_paginas_sin_excepcion():
    """Smoke test real (sin mocks): `App()` arma el sidebar y CADA pantalla
    -- incluidas Inicio/Acerca de, que quedan fuera del registro (SS4) --
    sin excepciones. Perezoso (SS4 del plan): una pagina recien se construye
    al visitarla, asi que hay que visitarlas todas a mano para ejercitarlas."""
    import customtkinter as ctk   # noqa: F401  (falla temprano y claro si falta la dep)
    from gui.app import App
    app = App()
    try:
        app.withdraw()   # no hace falta que la ventana se vea para este test
        ids = [p["id"] for p in app.registro] + ["inicio", "acerca_de"]
        for pid in ids:
            app.mostrar(pid)
            assert pid in app._frames, f"{pid}: no quedó construida tras mostrar()"
    finally:
        app.destroy()


def test_resumen_de_a23_y_sm_muestra_los_avisos():
    """Ronda 10, 2a pasada: los resumenes de A23 y SM decian "Listo" sin sus avisos
    (solo log y LEEME), y SM descartaba los del A03 (gemelo de §1.J.6, que solo arreglo
    Poblacion). Ahora los listan; sin avisos no inventan nada."""
    import pandas as pd
    from gui.paginas import a23 as pa23, sm as psm
    av = ("Seccion G (inasistentes cronicos)", "SUBCONTADO", "falta historial", "cargar mas")
    fer = pd.DataFrame({"RUN": ["A"]}); fer.attrs.update(seccion_g={}, avisos=[av])
    txt = pa23.resumen({"fer": fer, "salida": "x.xlsx", "mes": (2026, 7)})
    assert "SUBCONTADO" in txt and "falta historial" in txt, txt
    fer.attrs["avisos"] = []
    assert "aviso" not in pa23.resumen({"fer": fer, "salida": "x.xlsx", "mes": (2026, 7)})

    av03 = ("Tabla D.3", "SUBCONTADO", "1 aplicacion fuera de rango", "revisar")
    E = pd.DataFrame({"casilla": ["A04"]})
    E.attrs.update(tablas={"SM_Resumen": pd.DataFrame({"Casilla": ["A04"], "Total mes": [1]})},
                   avisos=[av])
    base = {"mes": (2026, 7), "solo_a03": False, "E": E, "n_tp": 0, "n_a03": 3,
            "por_inst_a03": {"PSC": 3}, "salida": "x.xlsx", "salida_a03": "y.xlsx",
            "avisos_a03": [av03]}
    txt = psm.resumen(base)
    assert "falta historial" in txt and "fuera de rango" in txt, txt
    txt = psm.resumen(dict(base, solo_a03=True))
    assert "fuera de rango" in txt, txt


def _main():
    pruebas = [v for k, v in sorted(globals().items())
              if k.startswith("test_") and callable(v)]
    fallos = 0
    for fn in pruebas:
        try:
            fn(); print(f"PASS  {fn.__name__}")
        except AssertionError as e:
            fallos += 1; print(f"FAIL  {fn.__name__} -> {e or 'assert'}")
        except Exception as e:   # noqa: BLE001
            import traceback
            fallos += 1; print(f"ERROR {fn.__name__} -> {type(e).__name__}: {e}")
            traceback.print_exc()
    print("-" * 50)
    print(f"{len(pruebas) - fallos}/{len(pruebas)} OK" + (f" ({fallos} problemas)" if fallos else ""))
    return 1 if fallos else 0


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    sys.exit(_main())
