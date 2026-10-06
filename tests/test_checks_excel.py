#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# ==========================================================================
# This code was generated with the assistance of Claude Opus 5.5 (Anthropic).
# The human author reviewed, modified, and integrated the code.
#
# Author: Simón Tobar — CESFAM Dr. Luis Ferrada Urzúa (APS, SSMC)
# SPDX-License-Identifier: GPL-3.0-or-later
# ==========================================================================
"""
Las macros de chequeo pre-envio (modulos/checks_excel/*.bas) leen celdas FIJAS de la
plantilla REM. Si MINSAL mueve una fila (plantilla nueva, o la 2027 de la 3.0.0), la
macro seguiria leyendo la celda vieja y diria "OK" sobre otra cosa: callado y errado.
Esto parsea cada .bas y exige que cada celda de total siga siendo una FORMULA en la
plantilla de refs_tablas/. Skill check-excel.
Correr: python tests/test_checks_excel.py
"""

import re
import sys
from functools import lru_cache
from pathlib import Path

import openpyxl
from openpyxl.utils import range_boundaries

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
import _aislar_cache   # noqa: E402,F401  (PRIMERO: nunca tocar el ~/.autorem real)

CHECKS = REPO / "modulos" / "checks_excel"
# REM_<serie>_... -> plantilla. Una serie nueva se agrega aca.
PLANTILLA = {"A": "SA_26*.xlsm", "P": "SP_26*.xlsm"}
# Totales que la plantilla deja como INPUT (sin formula): se verifican por la etiqueta
# de su fila, que es lo que se corre si MINSAL mueve algo.
MANUALES = {("A19a", "C110"): "salud mental", ("A19a", "C112"): "demencia"}
RE_ITEM = re.compile(r'"([^"|]+)\|[^"|]+\|([^"|]+)"')


@lru_cache(maxsize=None)
def _plantilla(serie):
    hits = sorted((REPO / "refs_tablas").glob(PLANTILLA[serie]))
    assert len(hits) == 1, f"serie {serie}: se esperaba UNA plantilla, hay {hits}"
    return openpyxl.load_workbook(hits[0])


def _macros():
    return sorted(CHECKS.glob("REM_*.bas"))


def _celdas(rango):
    for area in rango.split(","):
        c0, r0, c1, r1 = range_boundaries(area if ":" in area else f"{area}:{area}")
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                yield r, c


def test_hay_macros_y_cada_una_tiene_items():
    assert _macros(), "modulos/checks_excel/ sin macros"
    for bas in _macros():
        assert RE_ITEM.findall(bas.read_text(encoding="ascii")), f"{bas.name}: sin items"


def test_celdas_de_total_siguen_en_la_plantilla():
    malos = []
    for bas in _macros():
        wb = _plantilla(bas.stem.split("_")[1])
        for hoja, rango in RE_ITEM.findall(bas.read_text(encoding="ascii")):
            if hoja not in wb.sheetnames:
                malos.append(f"{bas.name}: no existe la hoja {hoja}")
                continue
            ws = wb[hoja]
            for r, c in _celdas(rango):
                cel = ws.cell(r, c)
                esperado = MANUALES.get((hoja, cel.coordinate))
                if esperado:
                    if esperado not in str(ws.cell(r, 2).value).lower():
                        malos.append(f"{bas.name}: {hoja}!B{r} ya no dice '{esperado}'")
                elif not str(cel.value).startswith("="):
                    malos.append(f"{bas.name}: {hoja}!{cel.coordinate} no es formula "
                                 f"({cel.value!r})")
    assert not malos, "\n".join(malos)


def _main():
    pruebas = [v for k, v in globals().items() if k.startswith("test_")]
    fallos = 0
    for fn in pruebas:
        try:
            fn(); print(f"OK    {fn.__name__}")
        except AssertionError as e:
            fallos += 1; print(f"FALLA {fn.__name__}\n{e}")
    return 1 if fallos else 0


if __name__ == "__main__":
    sys.exit(_main())
