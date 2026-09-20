import openpyxl, glob, os, sys
sys.stdout.reconfigure(encoding="utf-8")
for f in sorted(glob.glob("refs_tablas/*.xlsx")):
    if "CALCULADOR" in f: continue
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    for ws in wb.worksheets:
        print("=====", os.path.basename(f), "::", ws.title, ws.max_row, ws.max_column)
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if i > 14: print("  ..."); break
            vals = [(j, v) for j, v in enumerate(row) if v not in (None, "")]
            if vals: print(f"  r{i}:", " | ".join(f"{j}:{v!r}" for j, v in vals))
