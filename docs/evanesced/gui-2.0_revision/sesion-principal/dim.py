import zipfile, re, sys, shutil
from itertools import islice
sys.path.insert(0, r"E:\git\Autorem")
import openpyxl
from programas import formatos
from programas.formatos import MAX_FILAS_HEADER
S = sys.argv[1]
for src, name in [(r"refs_tablas\Formularios_RAYEN_csm_IRis.xlsx","iris"),(r"refs_tablas\Formulario_csm_reporte_Administrativo.xlsx","admin")]:
    z = zipfile.ZipFile(src)
    dims = {n: re.findall(rb'<dimension ref="[^"]*"/>', z.read(n)) for n in z.namelist() if n.startswith("xl/worksheets/sheet")}
    print(name, "dimension original:", dims)
    for label, repl in [("A1", b'<dimension ref="A1"/>'), ("ausente", b'')]:
        out = f"{S}/{name}_dim_{label}.xlsx"
        with zipfile.ZipFile(src) as zi, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
            for it in zi.infolist():
                data = zi.read(it.filename)
                if it.filename.startswith("xl/worksheets/sheet"):
                    data = re.sub(rb'<dimension ref="[^"]*"/>', repl, data)
                zo.writestr(it, data)
        wb = openpyxl.load_workbook(out, read_only=True, data_only=True)
        filas = list(islice(wb.active.iter_rows(values_only=True), MAX_FILAS_HEADER)); wb.close()
        cheap = formatos.detectar_eje_filas(filas)
        full = formatos.detectar_eje(openpyxl.load_workbook(out, data_only=True).active)
        from programas.rem_utils import leer_xlsx
        try: h, f = leer_xlsx(out); lx = (len(h), len(f))
        except Exception as e: lx = repr(e)[:80]
        print(f"  dim={label}: filas_leidas={len(filas)} ancho_fila0={len(filas[0]) if filas else 0} -> GUI(read_only)={cheap}  detectar_eje(full)={full}  leer_xlsx={lx}")
