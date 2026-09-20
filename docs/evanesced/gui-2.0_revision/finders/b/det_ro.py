import sys, zipfile, re
sys.path.insert(0, r"E:\git\Autorem")
import openpyxl
import programas.rem_saludmental as sm
files = [r"E:\git\Autorem\refs_tablas\Formularios_RAYEN_csm_IRis.xlsx",
         r"E:\git\Autorem\refs_tablas\Formulario_csm_reporte_Administrativo.xlsx"]
for f in files:
    z = zipfile.ZipFile(f)
    names = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
    for n in names:
        head = z.read(n)[:600].decode("utf-8", "ignore")
        m = re.search(r'<dimension ref="([^"]+)"', head)
        print(f.split("\\")[-1], n, "dimension:", m.group(1) if m else None)
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = wb.active
    print("  RO max_row/max_col:", ws.max_row, ws.max_column)
    try:
        print("  RO detect:", sm.detectar_formato(ws))
    except Exception as e:
        print("  RO detect EXC:", type(e).__name__, e)
    wb.close()
    wb = openpyxl.load_workbook(f)
    print("  full detect:", sm.detectar_formato(wb.active), wb.active.max_row, wb.active.max_column)
