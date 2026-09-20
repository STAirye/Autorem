import sys, zipfile, re, shutil, os
sys.path.insert(0, r"E:\git\Autorem")
import openpyxl
import programas.rem_saludmental as sm
S = sys.argv[1]
# IRIS-like formulario: 16 banner rows, header at 17, 3 data rows (fake RUT placeholder)
wb = openpyxl.Workbook(); ws = wb.active
for i in range(1, 17):
    ws.cell(row=i, column=1, value=f"banner {i}")
ws.cell(row=16, column=1, value=None)
hdr = ["NUMERO TIPO IDENTIFICACION", "SEXO", "AÑO APLICACIÓN FORMULARIO", "FECHA FORMULARIO", "INSTRUMENTO", "18.- ¿TIENE DEPRESION?"]
for j, h in enumerate(hdr, 1):
    ws.cell(row=17, column=j, value=h)
for r in range(18, 21):
    for j, v in enumerate(["11111111-1", "F", 30, "2026-08-01", "MEDICO", "SI"], 1):
        ws.cell(row=r, column=j, value=v)
wb.create_sheet("Hoja2"); wb.create_sheet("Hoja3")
base = os.path.join(S, "iris_ok.xlsx"); wb.save(base)

def rewrite(src, dst, fn):
    zin = zipfile.ZipFile(src); zout = zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED)
    for it in zin.infolist():
        data = zin.read(it.filename)
        if it.filename == "xl/worksheets/sheet1.xml":
            data = fn(data.decode("utf-8")).encode("utf-8")
        zout.writestr(it, data)
    zout.close(); zin.close()

rewrite(base, os.path.join(S, "iris_dimA1.xlsx"), lambda x: re.sub(r'<dimension ref="[^"]+"\s*/>', '<dimension ref="A1"/>', x))
rewrite(base, os.path.join(S, "iris_nodim.xlsx"), lambda x: re.sub(r'<dimension ref="[^"]+"\s*/>', '', x))

for name in ("iris_ok.xlsx", "iris_dimA1.xlsx", "iris_nodim.xlsx"):
    p = os.path.join(S, name)
    # exactly what gui/paginas/a05.py::_detectar.trabajo does
    try:
        w = openpyxl.load_workbook(p, read_only=True, data_only=True)
        cat = sm.detectar_formato(w.active); w.close()
    except Exception as e:
        cat = f"error_lectura ({type(e).__name__}: {e})"
    # what abrir_validado (old + new processing path) does
    try:
        wb2, ws2 = sm.abrir_validado(__import__("pathlib").Path(p), sm.PERFIL_IRIS)
        full = "accepted as IRIS"
    except Exception as e:
        full = f"{type(e).__name__}: {e}"
    print(f"{name:18s} GUI2.0 detection -> {cat!r:60s} | abrir_validado(IRIS) -> {full}")
