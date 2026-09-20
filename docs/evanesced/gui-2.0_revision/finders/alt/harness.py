# Harness (scratch, fuera del repo): alimenta exports SOLO-HEADER (0 filas de datos)
# a cada loader / entry point y reporta que pasa. Datos sinteticos, RUT 11111111-1.
import sys, os, traceback, shutil
from pathlib import Path
REPO = Path(r"E:\git\Autorem")
sys.path.insert(0, str(REPO))
OUT = Path(__file__).resolve().parent / "files"
OUT.mkdir(exist_ok=True)
import openpyxl
import pandas as pd

REFS = REPO / "refs_tablas"
quiet = lambda *a, **k: None
LOG = []
def logc(msg=""):
    LOG.append(str(msg))


def header_of(path, row_idx=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
    wb.close()
    return rows


def write(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def with_row(src, name, values):
    rows = header_of(src)
    hdr = rows[-1] if len([r for r in rows if any(v not in (None, "") for v in r)]) == 1 else None
    hdr = rows[0]
    row = [values.get(h, None) for h in hdr]
    return write(OUT / name, [hdr, row])


ADA_ROW = {"NUMERO TIPO IDENTIFICACION": "11111111-1", "FECHA ATENCION": "15/08/2026",
           "ACTIVIDADES": "Controles Salud Mental", "DIAGNOSTICOS": "F32 depresion",
           "INSTRUMENTO": "PSICOLOGO", "TIPO ATENCION": "CONTROL", "ATEN ID": "1",
           "PROFESIONAL ATENCION": "FUNC A", "SEXO": "Mujer", "AÑOS ATENCION": 30,
           "AÑOS": 30, "ALERTAS ADMINISTRATIVAS": "", "NACIONALIDAD": "CHILENA",
           "PUEBLO ORIGINARIO": "Ninguno", "FECHA DE NACIMIENTO": "01/01/1996",
           "ES IMIGRANTE": "NO", "FORMULARIOS CLINICOS": ""}
ada_1 = with_row(REFS / "ATENCIONESDIAGNOSTICOSACTIVIDADES_iris.xlsx", "ada_1row.xlsx", ADA_ROW)
ada_h = REFS / "ATENCIONESDIAGNOSTICOSACTIVIDADES_iris.xlsx"
grp_h = REFS / "Atenciones_Grupales_iris.xlsx"
insc_h = REFS / "Informe_Inscritos__Adscritos_heads.xlsx"
INSC_ROW = {"NUMERO TIPO IDENTIFICACION": "11111111-1", "TIPO IDENTIFICACION": "RUN",
            "SEXO": "Mujer", "GENERO": "Femenina", "FECHA DE NACIMIENTO": "01/01/1996",
            "EDAD AÑOS": 30, "SITUACION": "INSCRITO", "ESTADO": "ACTIVO", "SECTOR": "A",
            "ALERTAS ADMINISTRATIVAS": "", "PUEBLO INDIG": "", "NACIONALIDAD": "CHILENA"}
insc_1 = with_row(insc_h, "insc_1row.xlsx", INSC_ROW)
form_h = REFS / "Formularios_RAYEN_csm_IRis.xlsx"
FORM_ROW = {"NUMERO TIPO IDENTIFICACION": "11111111-1", "FECHA FORMULARIO": "10/08/2026",
            "INSTRUMENTO": "MEDICO", "SEXO": "Mujer", "AÑO APLICACIÓN FORMULARIO": 30,
            "18.- ¿ TIENE  DEPRESIÓN ?": "SI", "19.- ESTADO": "INGRESO"}
form_1 = with_row(form_h, "form_1row.xlsx", FORM_ROW)

otros_h = write(OUT / "otros_hdr.xlsx", [["NUMERO TIPO IDENTIFICACION", "FECHA ATENCION", "INSTRUMENTO",
                                           "SEXO", "FECHA DE NACIMIENTO", "1.- ¿PADECE DE ASMA BRONQUIAL?",
                                           "2.- ESTADO", "3.- FECHA DEL PROXIMO CONTROL"]])
estrat_h = write(OUT / "estrat_hdr.xlsx", [["RUT", "DV", "NOMBRE", "DETALLE DIAGNOSTICOS", "EDAD"]])
nsp_h = write(OUT / "nsp_hdr.xlsx", [["NUMERO TIPO IDENTIFICACION", "INSTRUMENTO", "TIPO ATENCION",
                                     "FECHA HORA CITA", "AÑOS"]])
multi_h = write(OUT / "multi_hdr.xlsx", [["ATEN ID", "FECHA", "Multiprofesional-1", "Multiprofesional-2"]])
maestro_h = write(OUT / "maestro_hdr.xlsx", [["ACTIVIDAD", "INSTRUMENTO ASOCIADO", "NUM REM", "NUM SECCION", "REM"]])
wrong_estrat = write(OUT / "wrong_estrat_hdr.xlsx", [["FOO", "BAR", "BAZ", "QUX", "QUUX"]])

import programas.rem_utils as ru
import programas.poblacion as pob
import modulos.rem_a23_respiratorio as a23
import modulos.rem_sm_actividades as smact
import modulos.rem_sm_trabajo_perdido as tp
import modulos.rem_sp_p6_poblacion as p6
import modulos.rem_sm_rescate_inasistentes as resc
import modulos.rem_a03_d3_instrumentos as screening
import programas.rem_saludmental as sm
import programas.estamentos as estam
DOT = {"funcionarios": {}, "omitidos": {}}
MES = (2026, 8)


def run(label, fn):
    LOG.clear()
    try:
        r = fn()
        extra = ""
        if isinstance(r, pd.DataFrame):
            extra = f" -> DataFrame len={len(r)}"
        elif isinstance(r, dict):
            extra = " -> " + ", ".join(f"{k}={(len(v) if hasattr(v, '__len__') and not isinstance(v, str) else v)}"
                                         for k, v in list(r.items())[:6])
        elif isinstance(r, tuple):
            extra = f" -> tuple {[type(x).__name__ for x in r]}"
        else:
            extra = f" -> {type(r).__name__} {r!r}"[:200]
        print(f"[NO ERROR ] {label}{extra}")
    except ru.ArchivoInvalido as e:
        print(f"[Archivo  ] {label}: ArchivoInvalido({e.categoria}) {str(e).splitlines()[0][:110]}")
    except Exception as e:
        tb = traceback.extract_tb(e.__traceback__)[-1]
        print(f"[CRYPTIC  ] {label}: {type(e).__name__}: {str(e)[:90]}  @ {Path(tb.filename).name}:{tb.lineno}")


run("rem_utils.cargar_atenciones(ADA header-only)", lambda: ru.cargar_atenciones(ada_h, log=quiet))
run("a23.procesar(ADA header-only)", lambda: a23.procesar(ada_h, mes=MES, log=quiet))
run("a23.cargar_otros(Otros header-only)", lambda: a23.cargar_otros(otros_h))
run("a23.procesar(ADA 1row, otros=header-only)", lambda: a23.procesar([ada_1], otros=[otros_h], mes=MES, log=quiet))
run("a23.cargar_estrat(Estrat header-only)", lambda: a23.cargar_estrat(estrat_h))
run("a23.cargar_estrat(WRONG file header-only)", lambda: a23.cargar_estrat(wrong_estrat))
run("a23.procesar(ADA 1row, inasistentes=NSP header-only)", lambda: a23.procesar([ada_1], inasistentes=[nsp_h], mes=MES, log=quiet))
run("smact.procesar(ADA header-only)", lambda: smact.procesar([ada_h], mes=MES, log=quiet, dotacion_tabla=DOT))
run("smact.procesar(ADA 1row, grupal header-only)", lambda: smact.procesar([ada_1], grupal=[grp_h], mes=MES, log=quiet, dotacion_tabla=DOT))
run("smact.procesar(ADA 1row, inscritos header-only)", lambda: smact.procesar([ada_1], inscritos=insc_h, mes=MES, log=quiet, dotacion_tabla=DOT))
run("smact.procesar(ADA 1row, multiprof header-only)", lambda: smact.procesar([ada_1], multiprofesional=multi_h, mes=MES, log=quiet, dotacion_tabla=DOT))
run("rem_utils.cargar_maestro(Maestro .xlsx header-only)", lambda: ru.cargar_maestro(maestro_h))
def _tp_maestro_vacio():
    E = tp.procesar([ada_1], maestro=maestro_h, mes=MES, log=logc)
    print("      tp log:", [l for l in LOG if "Maestro" in l], "| avisos:", E.attrs["avisos"],
          "| TP_Resumen:", E.attrs["tablas"]["TP_Resumen"].to_dict("records"))
    return E
run("tp.procesar(ADA 1row, maestro header-only)", _tp_maestro_vacio)
run("poblacion.cargar_inscritos(header-only)", lambda: pob.cargar_inscritos(insc_h, log=quiet))
run("poblacion.cargar_formulario_sm(header-only)", lambda: pob.cargar_formulario_sm([form_h], log=quiet))
run("poblacion.construir_poblacion(insc 1row, form header-only, ada 1row)",
    lambda: pob.construir_poblacion(insc_1, [form_h], [ada_1], mes=MES, log=quiet))
def _pob_ada_vacio():
    P = pob.construir_poblacion(insc_1, [form_1], [ada_h], mes=MES, log=logc)
    print("      avisos cobertura:", P.attrs["avisos"], "| Activo12m:", P["¿Activo 12m?"].tolist(),
          "| Ingresado:", P["¿Ingresado?"].tolist())
    return P
run("poblacion.construir_poblacion(insc 1row, form 1row, ADA header-only)", _pob_ada_vacio)
def _p6_ada_vacio():
    P = pob.construir_poblacion(insc_1, [form_1], [ada_h], mes=MES, log=quiet)
    return p6.construir_p6(P, log=quiet)
run("p6.construir_p6(P con ADA header-only)", _p6_ada_vacio)
run("resc.procesar(insc 1row, form 1row, ADA header-only)", lambda: resc.procesar(insc_1, [form_1], [ada_h], mes=MES, log=quiet))
run("A05 marcar_eventos IRIS header-only, mes=None (Archivo completo)",
    lambda: __import__("autorem")._correr_tareas(__import__("autorem").TAREAS, form_h, sm.PERFIL_IRIS, log=quiet, mes=None, carpeta=OUT))
run("A05 marcar_eventos ADMIN header-only, mes=None (Archivo completo)",
    lambda: __import__("autorem")._correr_tareas(__import__("autorem").TAREAS, REFS / "Formulario_csm_reporte_Administrativo.xlsx", sm.PERFIL_ADMIN, log=quiet, mes=None, carpeta=OUT))
run("A05 IRIS header-only, mes=(2026,8)",
    lambda: __import__("autorem")._correr_tareas(__import__("autorem").TAREAS, form_h, sm.PERFIL_IRIS, log=quiet, mes=MES, carpeta=OUT))
run("A03 procesar_unificado(GHQ-12 IRIS header-only)",
    lambda: screening.procesar_unificado({"GHQ-12": str(REFS / "goldberg_iris.xlsx")}, OUT / "a03.xlsx", log=quiet))
GHQ_ROW = {"NUMERO TIPO IDENTIFICACION": "11111111-1", "FECHA FORMULARIO": "10/08/2026", "SEXO": "Mujer",
           "AÑO APLICACIÓN FORMULARIO": 30, "1.- ESTADO": "Ingreso", "14.- PUNTAJE": 8,
           "15.- RESULTADO": "Indicativos de presencia de psicopatología", "FORMULARIO": "Cuestionario de Salud de Goldberg"}
ghq_1 = with_row(REFS / "goldberg_iris.xlsx", "ghq_1row.xlsx", GHQ_ROW)
run("A03 procesar_unificado(GHQ 1row + PSC-Y IRIS header-only)",
    lambda: screening.procesar_unificado({"GHQ-12": str(ghq_1), "PSC-Y": str(REFS / "pscy_10_14_iris.xlsx")}, OUT / "a03b.xlsx", log=quiet))
