import sys, zipfile, re, shutil, os
sys.path.insert(0, r"E:\git\Autorem")
import openpyxl
import programas.rem_saludmental as sm
SCR = os.path.dirname(os.path.abspath(__file__))

def dim_of(path):
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
        out = {}
        for n in names:
            x = z.read(n).decode("utf-8", "replace")
            m = re.search(r"<dimension[^>]*/>", x)
            out[n] = m.group(0) if m else None
        return out

def rewrite(src, dst, mode):
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                x = data.decode("utf-8")
                if mode == "nodim":
                    x = re.sub(r"<dimension[^>]*/>", "", x)
                elif mode == "a1":
                    x = re.sub(r"<dimension[^>]*/>", '<dimension ref="A1"/>', x)
                data = x.encode("utf-8")
            zout.writestr(item, data)

for f in ["Formularios_RAYEN_csm_IRis.xlsx", "Formulario_csm_reporte_Administrativo.xlsx"]:
    p = os.path.join(r"E:\git\Autorem\refs_tablas", f)
    print(f, dim_of(p))
    for mode in ["orig", "nodim", "a1"]:
        q = p if mode == "orig" else os.path.join(SCR, mode + "_" + f)
        if mode != "orig":
            rewrite(p, q, mode)
        try:
            wb = openpyxl.load_workbook(q, read_only=True, data_only=True)
            ws = wb.active
            print("  ", mode, "RO max_row=", ws.max_row, "max_col=", ws.max_column, end=" ")
            try:
                print("RO detect=", sm.detectar_formato(ws), end=" ")
            except Exception as e:
                print("RO detect EXC", type(e).__name__, e, end=" ")
            wb.close()
        except Exception as e:
            print("  ", mode, "load EXC", e, end=" ")
        wb2 = openpyxl.load_workbook(q)
        print("| normal detect=", sm.detectar_formato(wb2.active))
