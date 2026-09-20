import sys
from datetime import date
sys.path.insert(0, r"E:\git\Autorem"); sys.path.insert(0, r"E:\git\Autorem\tests")
import _aislar_cache  # noqa
import openpyxl
from programas.rem_utils import ArchivoInvalido
import programas.poblacion as pob
import test_rescate_inasistentes as TR
import test_a23 as TA
import modulos.rem_a23_respiratorio as a23
logs = []
L = lambda *a, **k: logs.append(" ".join(map(str, a)))

def strip_questions(p, keep=lambda v: True):
    wb = openpyxl.load_workbook(p); ws = wb.active
    for c in ws[17]:
        if isinstance(c.value, str) and c.value[:3].strip(".- ").isdigit() and not keep(c.value):
            c.value = "PREGUNTA RENOMBRADA " + c.value.split(".-", 1)[-1]
    wb.save(p); return p
ING = {"rut": "11111111-1", "fecha": date(2026, 8, 1), TR._Q[18] if hasattr(TR, "_Q") else "18.- ¿ TIENE  DEPRESIÓN ?": "SI"}
from test_sp_p6 import _Q
ING = {"rut": "11111111-1", "fecha": date(2026, 8, 1), _Q[18]: "SI", _Q[19]: "19.- INGRESO"}
for etq, keep in (("baseline", lambda v: True), ("sin columnas N.- (renombradas)", lambda v: False)):
    f = strip_questions(TR._mk_formulario([ING]), keep)
    try:
        P = pob.construir_poblacion(TR._mk_inscritos([{"rut": "11111111-1"}]), f,
                                    TR._mk_ada([TR._sm("11111111-1", date(2026, 8, 5)), TR._sm("11111111-1", date(2025, 7, 5))]),
                                    mes=(2026, 8), log=L)
        print(etq, "-> ingresados", int((P["¿Ingresado?"] == "SI").sum()), "avisos", [a[:2] for a in P.attrs["avisos"]])
    except ArchivoInvalido as e:
        print(etq, "-> ArchivoInvalido", e.categoria, str(e)[:120])

# A23 otros sin columna INSTRUMENTO
o = TA._mk_otros([{"RUN": "A", "FECHA": date(2025, 5, 1), "INSTR": "Médico", "ASMA_p": "Si", "ASMA_est": "Ingreso", "ASMA_grav": "Leve"}])
wb = openpyxl.load_workbook(o); ws = wb.active; ws.delete_cols(3); wb.save(o)
try:
    fer = a23.procesar(TA._mk([{"NUMERO TIPO IDENTIFICACION": "A", "FECHA ATENCION": date(2026, 7, 10), "INSTRUMENTO": "Médico",
                                "FECHA DE NACIMIENTO": date(1990, 1, 1)}]), otros=o, mes=(2026, 7), log=L)
    print("A23 otros sin col INSTRUMENTO -> SALA ASMA", int((fer["SALA ASMA"] == "SI").sum()), [a[:2] for a in fer.attrs["avisos"]])
except ArchivoInvalido as e:
    print("A23 otros sin col INSTRUMENTO -> ArchivoInvalido", e.categoria, str(e)[:120])
