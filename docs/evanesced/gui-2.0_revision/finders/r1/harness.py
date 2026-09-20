# -*- coding: utf-8 -*-
"""R1 empirical harness: synthetic fixtures (fake RUTs only) + stubbed customtkinter.
Writes ONLY under the scratchpad r1 dir. Never touches E:\\git\\Autorem."""
import os, sys, types, traceback, shutil
from pathlib import Path
from datetime import date, datetime

R1 = Path(__file__).resolve().parent
OUT = R1 / "out"
OUT.mkdir(exist_ok=True)
FAKEHOME = R1 / "fakehome"
FAKEHOME.mkdir(exist_ok=True)
os.environ["HOME"] = str(FAKEHOME)
os.environ["USERPROFILE"] = str(FAKEHOME)

REPO = Path(r"E:\git\Autorem")
REFS = REPO / "refs_tablas"
sys.path.insert(0, str(REPO))


# ---------------- customtkinter stub ----------------
class _Dummy:
    def __init__(self, *a, **k):
        pass

    def __getattr__(self, name):
        def _f(*a, **k):
            return _Dummy()
        return _f

    def __call__(self, *a, **k):
        return _Dummy()


def _install_ctk_stub():
    m = types.ModuleType("customtkinter")

    def __getattr__(name):
        return _Dummy
    m.__getattr__ = __getattr__
    m.CTkFrame = _Dummy
    m.CTk = _Dummy
    m.get_appearance_mode = lambda: "Light"
    sys.modules["customtkinter"] = m


_install_ctk_stub()

import openpyxl  # noqa: E402
from programas import dotacion, estamentos  # noqa: E402
from programas.rem_utils import norm, num_pregunta, dv_rut, ArchivoInvalido  # noqa: E402

# Redirect caches explicitly (belt and braces)
dotacion.RUTA_CACHE = FAKEHOME / ".autorem" / "dotacion.json"
estamentos.RUTA_CACHE = FAKEHOME / ".autorem" / "estamentos.json"
assert str(dotacion.RUTA_CACHE).startswith(str(R1))
assert str(estamentos.RUTA_CACHE).startswith(str(R1))


def rut(n):
    body = str(10000000 + n)
    return f"{body}-{dv_rut(body)}"


LOGS = []


def mklog(prefix=""):
    buf = []

    def log(msg=""):
        buf.append(str(msg))
    log.buf = buf
    return log


# ---------------- fixture builders ----------------
def ref_header(ref_name):
    """(header_row_values, rows_before_header(list of lists)) from a ref file."""
    wb = openpyxl.load_workbook(REFS / ref_name, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    # header = last non-empty row (refs are header-only)
    idx = max(i for i, r in enumerate(rows) if any(v not in (None, "") for v in r))
    return rows[idx], rows[:idx]


def _match_col(header, key):
    """key: int -> question number; str -> normalized exact header; ('subs', [..])."""
    hn = [norm(h) for h in header]
    if isinstance(key, int):
        for i, h in enumerate(header):
            if num_pregunta(h) == key:
                return i
        raise KeyError(key)
    if isinstance(key, tuple):
        toks = [norm(t) for t in key[1]]
        for i, h in enumerate(hn):
            if all(t in h for t in toks):
                return i
        raise KeyError(key)
    k = norm(key)
    for i, h in enumerate(hn):
        if h == k:
            return i
    raise KeyError(key)


def from_ref(ref_name, rows, out_name, extra_banner=None):
    """Copy the ref's banner+header and append `rows` (list of dict key->value)."""
    header, pre = ref_header(ref_name)
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in pre:
        ws.append(r)
    ws.append(header)
    for r in rows:
        line = [None] * len(header)
        for k, v in r.items():
            line[_match_col(header, k)] = v
        ws.append(line)
    p = OUT / out_name
    wb.save(p)
    return p


def simple(header, rows, out_name, pre=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in pre:
        ws.append(r)
    ws.append(header)
    for r in rows:
        ws.append([r.get(h) for h in header] if isinstance(r, dict) else list(r))
    p = OUT / out_name
    wb.save(p)
    return p


def copy_ref(ref_name, out_name=None):
    p = OUT / (out_name or ref_name)
    shutil.copy(REFS / ref_name, p)
    return p


# ADA (IRIS header from ref)
def ada_row(n, fecha, act="Controles Salud Mental", instr="Psicólogo(a)", prof="FUNC UNO",
            aten=None, diag="F32.1 Episodio depresivo moderado", tipo="Control"):
    return {"NUMERO TIPO IDENTIFICACION": rut(n), "TIPO IDENTIFICACION": "RUN",
            "ATEN ID": aten or f"AT{n}{fecha:%Y%m%d}{act[:4]}", "FECHA ATENCION": fecha,
            "ACTIVIDADES": act, "DIAGNOSTICOS": diag, "INSTRUMENTO": instr,
            "TIPO ATENCION": tipo, "PROFESIONAL ATENCION": prof, "SEXO": "Mujer",
            "AÑOS ATENCION": 30, "AÑOS": 30, "SECTOR": "Norte", "NACIONALIDAD": "Chilena",
            "ALERTAS ADMINISTRATIVAS": "", "ES IMIGRANTE": "NO", "PUEBLO ORIGINARIO": "Ninguno",
            "FECHA DE NACIMIENTO": date(1996, 1, 1), "FORMULARIOS CLINICOS": ""}


def mk_ada(rows, name):
    return from_ref("ATENCIONESDIAGNOSTICOSACTIVIDADES_iris.xlsx", rows, name)


# Grupal: ref header lacks an ASISTE column -> use test header (tests/test_sm_actividades.py)
GRP_HDR = ["NUMERO TIPO IDENTIFICACION", "FECHA ATENCION", "ACTIVIDADES",
           "ASISTE (SI/NO)", "SEXO", "EDAD", "INSTRUMENTO", "FUNCIONARIO PRESTADOR"]


def grp_row(n, fecha_txt, act="Intervencion Psicosocial Grupal", asiste="SI"):
    return {"NUMERO TIPO IDENTIFICACION": rut(n), "FECHA ATENCION": fecha_txt,
            "ACTIVIDADES": act, "ASISTE (SI/NO)": asiste, "SEXO": "Mujer",
            "EDAD": "30 años 1 meses 2 días", "INSTRUMENTO": "Psicólogo(a)",
            "FUNCIONARIO PRESTADOR": "FUNC UNO"}


def mk_grupal(rows, name):
    return simple(GRP_HDR, rows, name)


# Formulario SM IRIS (header from ref)
def form_row(n, fecha, dep="SI", estado="Ingreso", tipo="Depresión Moderada", instr="Médico"):
    return {"NUMERO TIPO IDENTIFICACION": rut(n), "FECHA FORMULARIO": fecha,
            "INSTRUMENTO": instr, ("subs", ("AÑO", "APLICACION")): 30, "SEXO": "Mujer",
            18: dep, 19: estado, 20: tipo}


def mk_form(rows, name):
    return from_ref("Formularios_RAYEN_csm_IRis.xlsx", rows, name)


# Inscritos (header from ref)
def insc_row(n, tipoid="RUN", estado="Activo", sexo="Mujer"):
    return {"TIPO IDENTIFICACION": tipoid, "NUMERO TIPO IDENTIFICACION": rut(n), "SEXO": sexo,
            "GENERO": "Femenina", "FECHA DE NACIMIENTO": date(1996, 1, 1), "EDAD AÑOS": 30,
            "NACIONALIDAD": "Chilena", "PUEBLO INDIG": "Ninguno", "ALERTAS ADMINISTRATIVAS": "",
            "SITUACION": "Inscrito", "ESTADO": estado, "SECTOR": "Norte",
            "FECHA PASIVACION": None, "MOTIVO PASIVACION": ""}


def mk_insc(rows, name):
    return from_ref("Informe_Inscritos__Adscritos_heads.xlsx", rows, name)


def run(label, fn):
    """Run fn(); print classification + details. Returns (result, exc)."""
    print("\n" + "=" * 100)
    print("SCENARIO:", label)
    try:
        res = fn()
        print("RESULT: SUCCESS (no exception)")
        return res, None
    except ArchivoInvalido as e:
        print(f"RESULT: ArchivoInvalido categoria={e.categoria!r}")
        print("  msg:", str(e).replace("\n", " | ")[:400])
        tb = traceback.extract_tb(e.__traceback__)
        print("  raised at:", f"{tb[-1].filename}:{tb[-1].lineno} in {tb[-1].name}")
        return None, e
    except Exception as e:  # noqa
        print(f"RESULT: CRASH {type(e).__name__}: {e}")
        traceback.print_exc(file=sys.stdout)
        return None, e
