import openpyxl, glob, os
for p in sorted(glob.glob(r"E:\git\Autorem\refs_tablas\*.xlsx")):
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    nonempty = [i for i, r in enumerate(rows) if any(v not in (None, "") for v in r)]
    full = [i for i, r in enumerate(rows) if sum(v not in (None, "") for v in r) > 3]
    print(os.path.basename(p), "rows:", len(rows), "nonempty:", len(nonempty), "rows>3cells:", len(full), "sheets:", wb.sheetnames)
    wb.close()
