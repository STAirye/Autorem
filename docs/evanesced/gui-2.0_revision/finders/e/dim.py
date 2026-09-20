import sys, zipfile, re
sys.path.insert(0, r"E:\git\Autorem")
import openpyxl
import programas.rem_saludmental as sm
for f in ["Formularios_RAYEN_csm_IRis.xlsx", "Formulario_csm_reporte_Administrativo.xlsx"]:
    p = r"E:\git\Autorem\refs_tablas\\" + f
    z = zipfile.ZipFile(p)
    names = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
    for n in names:
        x = z.read(n).decode("utf-8", "replace")
        m = re.search(r"<dimension[^>]*>", x)
        print(f, n, m.group(0) if m else None, "rows:", len(re.findall(r"<row ", x)))
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb.active
    print("  read_only max_row", ws.max_row, "detect:", sm.detectar_formato(ws))
    wb.close()
    wb = openpyxl.load_workbook(p)
    print("  full max_row", wb.active.max_row, "detect:", sm.detectar_formato(wb.active))
