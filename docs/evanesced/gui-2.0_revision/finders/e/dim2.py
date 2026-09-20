import sys, zipfile, re, shutil
sys.path.insert(0, r"E:\git\Autorem")
import openpyxl
import programas.rem_saludmental as sm

def patch(src, dst, fn):
    zin = zipfile.ZipFile(src)
    zout = zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "xl/worksheets/sheet1.xml":
            data = fn(data.decode("utf-8")).encode("utf-8")
        zout.writestr(item, data)
    zout.close()

for f in ["Formularios_RAYEN_csm_IRis.xlsx", "Formulario_csm_reporte_Administrativo.xlsx"]:
    src = r"E:\git\Autorem\refs_tablas\\" + f
    for tag, fn in [("nodim", lambda x: re.sub(r"<dimension[^>]*/>", "", x)),
                    ("dimA1", lambda x: re.sub(r"<dimension[^>]*/>", '<dimension ref="A1"/>', x))]:
        dst = f"{tag}_{f}"
        patch(src, dst, fn)
        try:
            wb = openpyxl.load_workbook(dst, read_only=True, data_only=True)
            ws = wb.active
            print(tag, f, "ro max_row", ws.max_row, end=" ")
            print("detect_ro:", sm.detectar_formato(ws), end=" ")
            wb.close()
        except Exception as e:
            print("detect_ro EXC:", type(e).__name__, e, end=" ")
        wb = openpyxl.load_workbook(dst)
        print("| full:", sm.detectar_formato(wb.active))
