"""Copia limpia de la tabla de actividades de epidemiologia: solo valores de celda.
Fuera: docProps con nombre de funcionaria, 17 controles ActiveX y la imagen EMF."""
import sys, zipfile
import openpyxl

src, dst = sys.argv[1], sys.argv[2]
wi = openpyxl.load_workbook(src, data_only=True)
wo = openpyxl.Workbook()
wo.remove(wo.active)
for ws in wi.worksheets:
    wn = wo.create_sheet(ws.title)
    for row in ws.iter_rows():
        for c in row:
            if c.value not in (None, ""):
                wn[c.coordinate] = c.value
    for k, d in ws.column_dimensions.items():
        if d.width:
            wn.column_dimensions[k].width = d.width
wo.properties.creator = None
wo.properties.lastModifiedBy = None
wo.save(dst)

a = [(ws.title, c.coordinate, c.value) for ws in wi.worksheets for r in ws.iter_rows() for c in r if c.value not in (None, "")]
wb = openpyxl.load_workbook(dst)
b = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r if c.value not in (None, "")]
assert a == b, "las celdas no calzan"
core = zipfile.ZipFile(dst).read("docProps/core.xml").decode()
print("celdas:", len(b), "| partes:", [n for n in zipfile.ZipFile(dst).namelist() if "activeX" in n or "media" in n or "drawing" in n])
print("creator en core.xml:", "creator>" in core and "<dc:creator></dc:creator>" not in core)
