"""Sonda de la API de python-calamine sobre un .xlsx SINTETICO: tipos por celda contra
openpyxl, excepciones ante un no-xlsx, y si libera el archivo."""
import datetime as dt, os, tempfile, zipfile
from pathlib import Path
import openpyxl, python_calamine as pc
from python_calamine import CalamineWorkbook

d = Path(tempfile.mkdtemp())
p = d / "tipos.xlsx"
wb = openpyxl.Workbook(); ws = wb.active
vals = [1, 2.0, 2.5, 10000000, True, "texto", "", "  ", "0123", "#N/A",
        dt.datetime(2026, 8, 1), dt.datetime(2026, 8, 1, 13, 5), dt.date(2026, 8, 1),
        dt.time(13, 5), "=1+1", None, "ñandú"]
ws.append(vals)
wb.create_sheet("vacia")
wb.save(p)

ro = openpyxl.load_workbook(p, read_only=True, data_only=True)
A = list(ro.active.iter_rows(values_only=True))[0]; ro.close()
cw = CalamineWorkbook.from_path(str(p))
print("version", getattr(pc, "__version__", "?"), "| sheets", cw.sheet_names)
print("attrs wb:", [a for a in dir(cw) if not a.startswith("_")])
sh = cw.get_sheet_by_index(0)
print("attrs sheet:", [a for a in dir(sh) if not a.startswith("_")])
B = sh.to_python(skip_empty_area=False)[0]
for a, b in zip(A, list(B) + [None] * (len(A) - len(B))):
    print(f"  openpyxl {type(a).__name__:<9}{a!r:<32} calamine {type(b).__name__:<9}{b!r}")
print("vacia ->", cw.get_sheet_by_index(1).to_python(skip_empty_area=False))
try:
    print("nrows=1 ->", len(sh.to_python(skip_empty_area=False, nrows=1)))
except TypeError as e:
    print("nrows no soportado:", e)

html = d / "falso.xlsx"; html.write_text("<html><body>x</body></html>")
for q in (html,):
    try:
        CalamineWorkbook.from_path(str(q))
    except Exception as e:
        print("no-xlsx ->", type(e).__module__, type(e).__name__, "| mro:",
              [c.__name__ for c in type(e).__mro__])
trunc = d / "trunc.xlsx"; trunc.write_bytes(p.read_bytes()[:3000])
try:
    CalamineWorkbook.from_path(str(trunc)).get_sheet_by_index(0).to_python()
except Exception as e:
    print("truncado ->", type(e).__name__)
try:
    os.replace(trunc, d / "movido.xlsx"); print("truncado liberado: OK")
except PermissionError:
    print("truncado BLOQUEADO tras la excepcion")
del cw, sh
try:
    os.replace(p, d / "movido2.xlsx"); print("tipos liberado tras del: OK")
except PermissionError:
    print("tipos BLOQUEADO")
