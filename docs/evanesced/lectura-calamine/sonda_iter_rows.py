"""iter_rows de python-calamine: arranca en A1 o en la primera celda con datos? Libro
sintetico con datos desde C3 y filas desparejas."""
import inspect, tempfile
from pathlib import Path
import openpyxl
from python_calamine import CalamineWorkbook

p = Path(tempfile.mkdtemp()) / "offset.xlsx"
wb = openpyxl.Workbook(); ws = wb.active
ws["C3"] = "banner"; ws["C5"] = "RUN"; ws["D5"] = "FECHA"; ws["E5"] = "x"
ws["C6"] = "r1"; ws["D6"] = 5
wb.save(p)
cw = CalamineWorkbook.from_path(str(p))
sh = cw.get_sheet_by_index(0)
print("start/end:", sh.start, sh.end, "height/width:", sh.height, sh.width,
      "total:", sh.total_height, sh.total_width)
print("to_python(skip_empty_area=False):", sh.to_python(skip_empty_area=False))
print("iter_rows():", list(sh.iter_rows()))
try:
    print("iter_rows sig:", inspect.signature(sh.iter_rows))
except (TypeError, ValueError) as e:
    print("sig no disponible:", e, "| doc:", (sh.iter_rows.__doc__ or "")[:300])
cw.close()
