"""A05 N+O sobre el formulario PSM real, SIN profiler, por fases (los mismos pasos de
autorem._correr_tareas). Ademas mide la alternativa de escritura: libro NUEVO en
write_only con solo las hojas de salida. Solo imprime tiempos."""
import sys, os, time, statistics, tempfile
from pathlib import Path
sys.path.insert(0, r"C:\Users\simon.tobar\Dr tobar\AutoREM")
import tests._aislar_cache  # noqa: F401
import openpyxl
import autorem
import programas.rem_saludmental as sm
from programas import cobertura
from programas.rem_utils import primeras_filas, verificar_hoja_unica

DM = Path(os.path.expanduser("~")) / "OneDrive - Ilustre Municipalidad de Maip\u00fa" / "Datos madre"
PSM = DM / "Variables 12m" / "PSM" / "2026.xlsx"
MES = (2026, 8)
TMP = Path(tempfile.mkdtemp(prefix="a05_"))
N = 3


def quiet(*a, **k):
    pass


perfil = sm.perfil_por_id(sm.detectar_formato_filas(primeras_filas(PSM, 30)))
fases = {k: [] for k in ("verificar_hoja_unica", "load_workbook+validar", "egresos",
                         "ingresos", "LEEME", "save (libro entero)",
                         "ALT: write_only solo salida", "TOTAL hoy")}
try:
    for i in range(N):
        t0 = time.perf_counter()
        t = time.perf_counter(); verificar_hoja_unica(PSM); fases["verificar_hoja_unica"].append(time.perf_counter() - t)
        t = time.perf_counter()
        wb = openpyxl.load_workbook(PSM)
        ws = wb.active
        perfil["validar"](ws)
        fases["load_workbook+validar"].append(time.perf_counter() - t)
        for tarea, k in zip(autorem.TAREAS, ("egresos", "ingresos")):
            t = time.perf_counter(); tarea["agregar"](wb, ws, perfil, log=quiet, mes=MES)
            fases[k].append(time.perf_counter() - t)
        t = time.perf_counter()
        cobertura.escribir_hoja(wb, [t_["id"] for t_ in autorem.TAREAS],
                                {"mes": MES, "archivos": [PSM.name]}, avisos=[])
        fases["LEEME"].append(time.perf_counter() - t)
        t = time.perf_counter(); wb.save(TMP / f"hoy_{i}.xlsx"); fases["save (libro entero)"].append(time.perf_counter() - t)
        fases["TOTAL hoy"].append(time.perf_counter() - t0)
        # alternativa: libro nuevo write_only con las hojas que NO son la original
        t = time.perf_counter()
        nuevo = openpyxl.Workbook(write_only=True)
        for h in wb.worksheets:
            if h is ws:
                continue
            dst = nuevo.create_sheet(h.title)
            for fila in h.iter_rows(values_only=True):
                dst.append(fila)
        nuevo.save(TMP / f"alt_{i}.xlsx")
        fases["ALT: write_only solo salida"].append(time.perf_counter() - t)
        wb.close()
    print(f"A05 N+O, {PSM.stat().st_size / 1e6:.1f} MB, mes {MES}, {N} corridas (mediana / min / max)")
    for k, v in fases.items():
        print(f"  {k:<30} {statistics.median(v):6.2f} s   {min(v):6.2f} / {max(v):6.2f}")
    print(f"  hojas en la salida de hoy: {len(openpyxl.load_workbook(TMP / 'hoy_0.xlsx', read_only=True).sheetnames)}")
finally:
    for p in TMP.glob("*"):
        p.unlink()
    TMP.rmdir()
