"""El `_hojas_calamine` nuevo (iter_rows + strings reusados) contra el de la rama antes
del cambio (to_python, copiado abajo tal cual), por repr y por identidad de TIPO, sobre
los exports reales y un sintetico con datos desde C3 + hojas vacias. Sin valores."""
import os, sys, tempfile
from pathlib import Path
sys.path.insert(0, os.getcwd())
import openpyxl
from python_calamine import CalamineWorkbook
from programas import rem_utils as ru


def viejo(entrada):   # la version de 16f589d
    wb = CalamineWorkbook.from_path(str(entrada))
    try:
        hojas = []
        for nombre in wb.sheet_names:
            crudas = wb.get_sheet_by_name(nombre).to_python(skip_empty_area=False)
            ancho = max((len(f) for f in crudas), default=0)
            hojas.append((nombre, [tuple(map(ru._celda_openpyxl, f)) + ("",) * (ancho - len(f))
                                   for f in crudas]))
        return hojas
    finally:
        wb.close()


sint = Path(tempfile.mkdtemp()) / "offset.xlsx"
wb = openpyxl.Workbook(); ws = wb.active
ws["C3"] = "banner"; ws["C5"] = "RUN"; ws["D5"] = "FECHA"; ws["E5"] = 2.5
ws["C6"] = "r1"; ws["D6"] = 5; ws["F9"] = "lejos"
wb.create_sheet("vacia1"); wb.create_sheet("vacia2")
wb.save(sint)

DM = Path(os.path.expanduser("~")) / "OneDrive - Ilustre Municipalidad de Maipú" / "Datos madre"
for nombre, p in (("sintetico C3", sint),
                  ("ADA", DM / "Variables 12m" / "Atenciones" / "2026.xlsx"),
                  ("Inscritos", DM / "Población" / "Informe_Inscritos__Adscritos_.xlsx"),
                  ("PSM", DM / "Variables 12m" / "PSM" / "2026.xlsx")):
    a, b = viejo(p), ru._hojas_calamine(p)
    forma = [(n, len(f), max((len(x) for x in f), default=0)) for n, f in a] == \
            [(n, len(f), max((len(x) for x in f), default=0)) for n, f in b]
    dif = sum(1 for (_, fa), (_, fb) in zip(a, b) for ra, rb in zip(fa, fb)
              for x, y in zip(ra, rb) if repr(x) != repr(y) or type(x) is not type(y))
    print(f"  {nombre:<13} misma forma (hojas/filas/ancho): {forma}   celdas distintas: {dif}")
