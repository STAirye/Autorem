"""Costo del barrido de filas de `marcar_eventos` (lo que se paga UNA VEZ POR TAREA
del A05: ingresos y egresos corren los dos sobre el MISMO ws)."""
import sys, time
sys.path.insert(0, r"C:\Users\simon.tobar\Dr tobar\AutoREM")
import openpyxl
from programas.rem_utils import norm, mes_de_celda, edad_anios

ruta = r"C:\Users\SIMON~1.TOB\AppData\Local\Temp\bench_form.xlsx"
import os
ruta = os.path.join(os.environ.get("TEMP", "/tmp"), "bench_form.xlsx")
NCOLS, HDR = 100, 16

wb = openpyxl.load_workbook(ruta, data_only=True)
ws = wb.active


def barrido():
    """Lo mismo que hace marcar_eventos por fila, sin el matching de tokens."""
    n = 0
    for r in range(HDR + 1, ws.max_row + 1):
        fila = [ws.cell(row=r, column=c).value for c in range(1, NCOLS + 1)]
        fila_n = [norm(v) for v in fila]          # <-- TODAS las columnas
        rut = fila[0]
        if not str(rut or "").strip():
            continue
        mes_de_celda(fila[1])
        edad_anios(fila[2])
        n += len(fila_n)
    return n


for i in range(3):
    t = time.perf_counter()
    n = barrido()
    print(f"   barrido completo #{i+1}: {time.perf_counter()-t:.3f} s  ({n} celdas normalizadas)")
wb.close()
print("   -> el A05 con las DOS tareas marcadas paga este barrido DOS veces")
