#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# ==========================================================================
# This code was generated with the assistance of Claude Opus 5.5 (Anthropic).
# The human author reviewed, modified, and integrated the code.
#
# Author: Simon Tobar - CESFAM Dr. Luis Ferrada Urzua (APS, SSMC)
# Copyright (C) 2026 Simon Tobar
# SPDX-License-Identifier: GPL-3.0-or-later
# (Sin 'Version:': herramienta de la skill, fuera del codigo distribuible.)
# ==========================================================================
"""
inventario.py - borrador de inventario de casillas de un REM llenado a mano.

Lee una planilla REM (SA_26 / SP_26 .xlsm) que un equipo lleno a mano y saca,
por seccion: el mapa de columnas de input, las casillas llenas, las columnas de
formula, las alertas de validacion que la planilla ya disparo, las notas del
REM Comentado (cruzadas por ETIQUETA, no por coordenada) y las actividades del
Maestro que apuntan a esa seccion. Escribe un borrador .md para que Claude lo
convierta en docs/<programa>_casillas_<serie>.md (skill inventario-rem).

Input = casilla DESBLOQUEADA (las hojas REM vienen protegidas y solo el input
queda sin candado). La planilla nunca se copia: el borrador va al scratchpad.

Uso:
  python .claude/skills/inventario-rem/inventario.py REM.xlsm --salida borrador.md
      [--hojas A05,A26] [--secciones "A05:J,V;A33:*"] [--filas REGEX]
      [--comentado refs_tablas/REM_Comentado_Serie_A_2026_15-04-2026.xlsx]
      [--maestro]
"""
import argparse
import re
import sys
import warnings
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as L

RAIZ = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(RAIZ))
from programas.rem_utils import norm  # noqa: E402

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
RE_SEC = re.compile(r"^SECCI\w*N\s+([A-Z][0-9A-Z.]*?)\s*[:.]?\s", re.I)


def es_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) and v != 0


def es_txt(v):
    return isinstance(v, str) and v.strip() != "" and not v.startswith("=")


def es_formula(v):
    return isinstance(v, str) and v.startswith("=")


def limpio(s):
    return re.sub(r"\s+", " ", str(s).replace(chr(0x2060), "")).strip()  # WORD JOINER de RAYEN


class Hoja:
    """Una hoja REM con celdas combinadas resueltas y sus secciones."""

    def __init__(self, ws, ws_valores=None):
        self.ws, self.wv = ws, ws_valores
        self.mm = {}
        for rg in ws.merged_cells.ranges:
            for r in range(rg.min_row, rg.max_row + 1):
                for c in range(rg.min_col, rg.max_col + 1):
                    self.mm[(r, c)] = (rg.min_row, rg.min_col)
        ini = [r for r in range(1, ws.max_row + 1)
               if RE_SEC.match(norm(ws.cell(r, 1).value) + " ")]
        ini.append(ws.max_row + 1)
        self.secciones = []
        for a, b in zip(ini, ini[1:]):
            titulo = limpio(ws.cell(a, 1).value)
            cod = RE_SEC.match(norm(titulo) + " ").group(1).rstrip(".")
            self.secciones.append((cod, titulo, a, b))

    def val(self, r, c):
        r0, c0 = self.mm.get((r, c), (r, c))
        return self.ws.cell(r0, c0).value

    def es_input(self, r, c):
        return (r, c) not in self.mm and not self.ws.cell(r, c).protection.locked

    def filas_datos(self, a, b):
        maxc = self.ws.max_column
        d = {r: [c for c in range(1, maxc + 1) if self.es_input(r, c)] for r in range(a + 1, b)}
        return {r: cs for r, cs in d.items() if cs}

    def etiqueta(self, r, hasta):
        lab = []
        for c in range(1, hasta):
            t = self.val(r, c)
            if es_txt(t) and limpio(t) not in lab:
                lab.append(limpio(t))
        return lab

    def encabezado(self, r, c, a, datos):
        """Texto de encabezado de la columna c para la fila r: el bloque de filas
        no-dato contiguo mas cercano por arriba (una seccion puede tener varios)."""
        k = r - 1
        while k > a and k in datos:
            k -= 1
        h = []
        while k > a and k not in datos:
            t = self.val(k, c)
            if es_txt(t) and limpio(t) not in h:
                h.insert(0, limpio(t))
            k -= 1
        return " > ".join(h)


def notas_comentado(hc, cod):
    """{clave_fila: [(fila, [(coord, encabezado, texto)])]} y notas de columna."""
    sec = next((s for s in hc.secciones if s[0] == cod), None)
    if not sec:
        return None, [], None
    _, titulo, a, b = sec
    datos = hc.filas_datos(a, b)
    por_fila, de_columna = {}, []
    for r in range(a, b):
        for cell in hc.ws[r]:
            if not cell.comment:
                continue
            texto = re.sub(r"\n\s*\n+", "\n", cell.comment.text.strip())
            if r in datos:
                cab = hc.encabezado(r, cell.column, a, datos)
                lab = " . ".join(hc.etiqueta(r, datos[r][0]))
                por_fila.setdefault(r, (lab, []))[1].append((cell.coordinate, cab, texto))
            else:
                de_columna.append((cell.coordinate, limpio(cell.value or ""), texto))
    return por_fila, de_columna, (titulo, a, b)


def clave(s):
    return re.sub(r"[^A-Z0-9]+", " ", norm(s)).strip()


def emparejar(etiquetas_rem, filas_com, umbral=0.75):
    """Empareja filas del REM con filas del Comentado por ETIQUETA, en orden y
    sin repetir (hay etiquetas duplicadas, p.ej. 'Egreso por traslado' en A05 V
    para la persona y para el cuidador). Match aproximado: los textos difieren
    entre versiones ('Cuidador/a' vs 'CUIDADOR (A)'). Devuelve {fila_rem: fila_com}."""
    from difflib import SequenceMatcher
    libres = sorted(filas_com)
    par = {}
    for r, lab in etiquetas_rem:
        mejor = max(((SequenceMatcher(None, clave(lab), clave(filas_com[k][0])).ratio(), -k, k)
                     for k in libres), default=None)
        if mejor and mejor[0] >= umbral:
            par[r] = mejor[2]
            libres = [k for k in libres if k > mejor[2]]
    return par


def agrupar_notas(notas):
    """Junta columnas con el mismo texto de nota (el Comentado repite la regla
    casilla por casilla, cambiando solo el estamento o la causal)."""
    grupos = {}
    for coord, cab, texto in notas:
        lineas = [limpio(x) for x in texto.split("\n") if limpio(x)]
        tronco, cola = "\n".join(lineas[:-1]), lineas[-1] if lineas else ""
        grupos.setdefault(tronco, []).append((coord + (f" ({cab.rpartition(' > ')[2]})" if cab else ""), cola))
    res = {}
    for tronco, items in grupos.items():
        if len(items) == 1 or len({c for _, c in items}) == 1:
            res[(tronco + "\n" + items[0][1]).strip()] = [co for co, _ in items]
        else:
            res[tronco + "\n" + "\n".join(f"-> {co}: {c}" for co, c in items)] = [co for co, _ in items]
    return res


_MAESTRO = []


def maestro_por_seccion(hoja, cod):
    if not _MAESTRO:
        from programas.catalogos import maestro_slim
        from programas.rem_utils import cargar_maestro
        _MAESTRO.append(cargar_maestro(maestro_slim()))
    m = _MAESTRO[0]
    sec = cod.replace(".", "")
    # SECCION puede ser compuesta: 'A,B,L', 'A2,A3,F', 'A.6/B.6'
    d = m[m["NUMREM"].astype(str) == f"REM-{hoja}"]
    ok = d["SECCION"].astype(str).map(
        lambda s: sec in s.replace(".", "").replace("/", ",").replace(" ", "").split(","))
    return sorted(set(d.loc[ok, "ACT"].astype(str).str.strip()))


def inventariar(args):
    wb = openpyxl.load_workbook(args.rem)
    wv = openpyxl.load_workbook(args.rem, data_only=True)
    wc = openpyxl.load_workbook(args.comentado) if args.comentado else None
    filtro = re.compile(args.filas, re.I) if args.filas else None
    pedidas = {}
    for trozo in filter(None, (args.secciones or "").split(";")):
        h, _, ss = trozo.partition(":")
        pedidas[h.strip()] = [s.strip().rstrip(":") for s in ss.split(",")]
    hojas = pedidas.keys() if pedidas else (args.hojas.split(",") if args.hojas else wb.sheetnames)
    out = [f"# Borrador de inventario: {Path(args.rem).name}\n",
           "Generado por .claude/skills/inventario-rem/inventario.py. NO versionar: "
           "trae valores reales; pasarlo a docs/ via la skill.\n"]
    for nombre in hojas:
        if nombre not in wb.sheetnames:
            raise SystemExit(f"ERROR: la hoja {nombre} no existe en {args.rem}")
        h = Hoja(wb[nombre], wv[nombre])
        hc = Hoja(wc[nombre]) if wc and nombre in wc.sheetnames else None
        for cod, titulo, a, b in h.secciones:
            datos = h.filas_datos(a, b)
            con_valor = {r for r, cs in datos.items() if any(es_num(h.ws.cell(r, c).value) for c in cs)}
            llenas = bool(con_valor)
            quiere = pedidas.get(nombre)
            if quiere is not None:
                if "*" not in quiere and cod not in quiere:
                    continue
            elif not llenas:
                continue
            out.append(f"\n## {nombre} - {titulo}\n\nFilas {a}-{b - 1}. "
                       f"{'Con casillas llenas.' if llenas else 'SIN casillas llenas.'}\n")
            bloques = {}
            for r, cs in datos.items():
                bloques.setdefault(tuple(cs), []).append(r)
            for cs, filas in bloques.items():
                r0 = filas[0]
                cab = {c: h.encabezado(r0, c, a, datos) for c in cs}
                form = sorted({L(c) for r in filas for c in range(1, h.ws.max_column + 1)
                               if es_formula(h.ws.cell(r, c).value)},
                              key=lambda x: (len(x), x))
                out.append(f"**Columnas de input** (bloque desde fila {r0}; formulas en "
                           f"{', '.join(form) or 'ninguna'}):\n")
                grupos = {}
                for c, t in cab.items():
                    pre, _, hoja_ = t.rpartition(" > ")
                    grupos.setdefault(pre, []).append(f"`{L(c)}` {hoja_ or t or '?'}")
                out += [f"- {p or '(sin encabezado)'}: " + " . ".join(i) for p, i in grupos.items()]
                out.append("\n| Fila | Etiqueta | Casillas llenas | Suma |\n|---|---|---|---|")
                for r in filas:
                    lab = " . ".join(h.etiqueta(r, cs[0]))
                    llen = [(c, h.ws.cell(r, c).value) for c in cs if es_num(h.ws.cell(r, c).value)]
                    if filtro and not llen and not filtro.search(lab):
                        continue
                    txt = " . ".join(f"`{L(c)}{r}` {cab[c].rpartition(' > ')[2]}: **{v:g}**"
                                     for c, v in llen) or "-"
                    out.append(f"| {r} | {lab} | {txt} | {sum(v for _, v in llen):g} |")
                out.append("")
            alertas = [(c.coordinate, limpio(h.wv[c.coordinate].value)) for r in datos
                       for c in h.ws[r] if es_formula(c.value) and "&" in c.value
                       and h.wv[c.coordinate].value not in (None, "")]
            if alertas:
                out.append("**Alertas de validacion ya disparadas** (valor cacheado; puede estar viejo):\n")
                out += [f"- `{co}`: {t[:220]}" for co, t in alertas]
                out.append("")
            if hc:
                por_fila, de_col, info = notas_comentado(hc, cod)
                if info is None:
                    out.append(f"**Comentado:** no hay seccion {cod} en la hoja {nombre}.\n")
                else:
                    out.append(f"**Comentado** ({info[0]}, filas {info[1]}-{info[2] - 1} alli):\n")
                    for co, cab_, t in de_col:
                        out.append(f"- Columna `{co}` {cab_}: {t}")
                    visibles = [(r, " . ".join(h.etiqueta(r, datos[r][0]))) for r in datos]
                    par = emparejar(visibles, por_fila)
                    for r, lab in visibles:
                        if r not in par or filtro and r not in con_valor and not filtro.search(lab):
                            continue
                        rc = par[r]
                        out.append(f"\n- **Fila {r}** = fila {rc} del Comentado ({por_fila[rc][0]}):")
                        for texto, coords in agrupar_notas(por_fila[rc][1]).items():
                            out.append(f"  - [{', '.join(coords)}]\n    " + texto.replace("\n", "\n    "))
                    sin_par = [f"{k} ({v[0]})" for k, v in por_fila.items() if k not in par.values()
                               and not (filtro and not filtro.search(v[0]))]
                    if sin_par:
                        out.append(f"\n- AVISO: filas del Comentado sin par en el REM: {sin_par}")
                    sin_nota = [str(r) for r, lab in visibles if r not in par
                                and not (filtro and r not in con_valor and not filtro.search(lab))]
                    if sin_nota:
                        out.append(f"- AVISO: filas del REM sin nota en el Comentado: {', '.join(sin_nota)}")
                    out.append("")
            if args.maestro:
                acts = maestro_por_seccion(nombre, cod)
                out.append(f"**Maestro** (REM-{nombre} / seccion {cod}): {len(acts)} actividades.\n")
                out += [f"- {x}" for x in acts]
                out.append("")
    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("rem")
    ap.add_argument("--salida", required=True)
    ap.add_argument("--hojas")
    ap.add_argument("--secciones")
    ap.add_argument("--filas")
    ap.add_argument("--comentado")
    ap.add_argument("--maestro", action="store_true")
    args = ap.parse_args()
    texto = inventariar(args)
    Path(args.salida).write_text(texto, encoding="utf-8")
    print(f"OK - borrador escrito en {args.salida} ({texto.count(chr(10))} lineas)")


if __name__ == "__main__":
    main()
