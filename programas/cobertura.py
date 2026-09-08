#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# ==========================================================================
# This code was generated with the assistance of Claude Sonnet 5 (Anthropic).
# The human author reviewed, modified, and integrated the code.
#
# Author: Simón Tobar — CESFAM Dr. Luis Ferrada Urzúa (APS, SSMC)
# Copyright (C) 2026 Simón Tobar
# SPDX-License-Identifier: GPL-3.0-or-later
# Version: 1.9.1
#
# This program is free software: you can redistribute it and/or modify it
# under the terms of the GNU General Public License as published by the
# Free Software Foundation, either version 3 of the License, or (at your
# option) any later version. Distributed WITHOUT ANY WARRANTY. See the GNU
# General Public License for more details: <https://www.gnu.org/licenses/>.
# ==========================================================================
"""
programas/cobertura.py — hoja «LEEME»: qué NO cubre autoREM (ver docs/hoja_cobertura_plan.md).

El exe corre sin errores y la planilla sale, pero nada en ella dice qué casillas
del REM de ese módulo NO quedaron llenas (algunas son manuales por diseño; otras
se degradan según qué fuentes opcionales se cargaron ESTA corrida). Ese aviso hoy
vive en docstrings, en CLAUDE.md (no se distribuye con el exe) o en el log de la
GUI (se pierde al cerrar la ventana) — exactamente la clase de "número plausible
pero callado" que la regla del proyecto "fallar ruidoso" existe para evitar.

Fuente ÚNICA declarativa: el catálogo vive acá, no repetido en cada módulo.
`tests/test_cobertura.py` es el guardarraíl anti-olvido (§7 del plan): un módulo
nuevo en `modulos/` sin entrada acá rompe el test.

Dos capas por módulo:
  - ESTRUCTURAL (`COBERTURA[id]['no_cubre']`): lo que el módulo NUNCA cubre
    (manual / fuera de alcance / omitido a propósito / pendiente / en validación /
    sin registro). Siempre igual, no depende de esta corrida.
  - DINÁMICA (`avisos`): lo que quedó a medias por los INPUTS de esta corrida
    (sin reporte opcional cargado, etc.) — la aportan los módulos vía
    `.attrs['avisos']` en su DataFrame de salida (patrón ya usado para
    `.attrs['tablas']`).

API:
  entradas(modulo_id, avisos=())        -> dict {rem, cubre, no_cubre, avisos}
  escribir_hoja(wb, modulo_id, ...)     -> crea la hoja LEEME como PRIMERA del wb
"""

import datetime

from programas.rem_utils import VERSION

# -- Categorías (el "por qué" no está lleno). El texto es lo que ve el usuario. --
MANUAL = "MANUAL"                 # no hay reporte/formulario en RAYEN: se llena a mano
FUERA = "FUERA DE ALCANCE"        # otro REM, u otro módulo del exe
OMITIDO = "OMITIDO"               # existe pero este centro no lo usa
PENDIENTE = "PENDIENTE"           # se va a implementar, hoy sale 0
VALIDACION = "EN VALIDACION"      # implementado pero los números NO están cerrados
SIN_REGISTRO = "SIN REGISTRO"     # el 0 es CORRECTO: nadie usa el canal de registro


# ======================================================================
# Catálogo (§3 del plan). Cada entrada: (casilla, categoria, motivo, que_hacer).
# ======================================================================
COBERTURA = {
    "sm_actividades": {
        "rem": "SA_26 - Salud Mental (A04 / A06 / A19a / A26 / A27 / A32)",
        "cubre": ["A04-A24", "A06-A.1", "A19a-A.3", "A26-A", "A27", "A32-F"],
        "no_cubre": [
            ("A06-A.2 Consultorias de Salud Mental", MANUAL,
             "No hay reporte ni formulario en RAYEN que las registre",
             "Contarlas a mano y pegarlas en el SA_26"),
            ("A05 ingresos / egresos", FUERA,
             "Los cubre otro modulo del exe", "Usar la pestana A05"),
            ("A03-H Tamizaje (PSC-17, PHQ-9...)", FUERA,
             "El exe cubre A03-D.3 (instrumentos de personas ya INGRESADAS)", ""),
            ("Espacios Amigables - Familias en Riesgo", OMITIDO,
             "No se usan en este centro (rem_sm_actividades.py)",
             "Si el centro empieza a usarlos, hay que implementarlos"),
            ("Campana de Invierno (col. demografica)", OMITIDO,
             "Write-protected en la hoja SM del template", ""),
            ("Demografia de las actividades grupales", PENDIENTE,
             "El reporte 'Atenciones Grupales' no trae demografia -> A06 psicosocial "
             "/ A19a grupal / A27 salen con demografia en 0",
             "Si se necesita, cruzar a mano contra el padron"),
            ("Control SM a paciente SENAME", OMITIDO,
             "Excluido a proposito: SENAME hace su propio REM", ""),
            ("A27 y A32-F2", SIN_REGISTRO,
             "La actividad existe en RAYEN pero no se usa (se agrego hace poco). El "
             "trabajo se registra en la ficha, en indicaciones, que no tributa -> 0 "
             "actividades realizadas. El filtro esta implementado pero nunca se pudo "
             "validar contra datos reales",
             "El 0 es correcto: no lo llenes a mano asumiendo que falta. Si sabes que "
             "la actividad se hizo, se registro en un canal que el REM no ve"),
        ],
    },
    "sm_trabajo_perdido": {
        "rem": "Auditoria (no tributa a ninguna casilla del REM)",
        "cubre": [],
        "no_cubre": [
            ("(todas)", FUERA,
             "Este reporte no tributa a ninguna casilla del REM: es auditoria de "
             "trabajo mal registrado", "No copiar nada de aqui al SA_26"),
            ("A26-A1 VDI del PADDS (dependencia severa)", PENDIENTE,
             "Se excluyen del universo SM a proposito (EXCLUIR_SMISH) porque no son "
             "SM - pero hoy no las tabula ningun modulo",
             "Contarlas a mano hasta que exista rem_a26_domiciliaria"),
        ],
    },
    "a23_respiratorio": {
        "rem": "SA_26 - A23 Respiratorio",
        "cubre": ["A23-D", "A23-E", "A23-F", "A23-G", "A23-H", "A23-I (parcial)",
                  "A23-M.1", "A23-N (parcial)"],
        "no_cubre": [
            ("Secciones B - C - J - K - L - M.2 - P - Q", FUERA,
             "Fuera de alcance del modulo", "A mano"),
            ("Seccion A (ingreso a sala)", PENDIENTE,
             "Semantica distinta: 'ingreso a sala' != 'tuvo dx' -> sale 0", "A mano"),
            ("Seccion I espirometria basal / post BD", PENDIENTE,
             "Hoy hay 1 solo indicador, el REM pide dos", "A mano"),
            ("Seccion O (EPOC forma A/B)", PENDIENTE,
             "Falta la forma A/B -> sale 0", "A mano"),
            ("Formulario 'Otros Cronicos' en formato Administrativo", PENDIENTE,
             "Solo se lee el formato IRIS", "Descargar el formulario en IRIS"),
            ("Ira Alta / Bronquitis / EPOC exacerbado", PENDIENTE,
             "Si cargaste el Monitoreo Administrativo en vez del reporte IRIS, estos "
             "indicadores salen en 0: el monitoreo trae el diagnostico en texto sin "
             "codigo ICD. autoREM todavia no puede detectar cual de los dos cargaste "
             "(depende de formatos.py fase 2)",
             "Verifica que cargaste el export IRIS completo, no el Monitoreo Administrativo"),
        ],
    },
    "a05_o_egresos": {
        "rem": "SA_26 - A05 Salud Mental (Egresos / Ingresos)",
        "cubre": ["A05 Egresos"],
        "no_cubre": [
            ("Egresos por Otras Causas (clasificacion)", MANUAL,
             "Requiere decision clinica caso a caso (abandono vs clinica). El modulo "
             "los flaggea, no los clasifica",
             "Revisar la lista de RUT del detalle y clasificar"),
            ("Epilepsia (pregunta 75)", FUERA,
             "Tributa al REM de adulto, no a SM (EXCLUIR_PATOLOGIA)", ""),
            ("Programas de rehabilitacion / acompanamiento (77, 79, 81)", FUERA,
             "No son diagnosticos SM", ""),
        ],
    },
    "a03_d3_instrumentos": {
        "rem": "SA_26 - A03 D.3 (instrumentos)",
        "cubre": ["A03-D.3"],
        "no_cubre": [
            ("A03-H Tamizaje", FUERA,
             "El modulo cubre D.3 (personas INGRESADAS al PSM)", "A mano"),
            ("Aplicaciones 'Sin riesgo' (bajo el corte)", OMITIDO,
             "Van al detalle auditable pero no al D.3 (por diseno)", ""),
            ("Conteos agregados por rango etario", PENDIENTE,
             "Planificado para A03 D.3 v2", ""),
        ],
    },
    "sp_p6_poblacion": {
        "rem": "SP_26 - P6-A.1 Poblacion en control PSM",
        "cubre": ["SP-P6-A.1"],
        "no_cubre": [
            ("Plan de Cuidado Integral (col. AV)", SIN_REGISTRO,
             "No va vacia ni a mano: la llena una regla operativa WIP (total de la "
             "fila en GES depresion/Alzheimer y en factores de riesgo; 0 en el "
             "resto). La fuente existe (actividad 'Plan Cuidado Integral Elaborado' "
             "del Maestro) pero el registro es ~0: el PIC se escribe en indicaciones "
             "de la ficha. Regla completa en SP_P6_poblacion_plan.md 5.4.2",
             "Revisar AV24 a mano: si AV24 > C24 hay comorbilidad depresion+demencia "
             "contada dos veces"),
            ("Toda la grilla", VALIDACION,
             "Modulo en validacion: el filtro Ingresado da 2972 contra 2226 del "
             "PowerBI (ver SP_P6_poblacion_plan.md 9)",
             "Contrastar contra el conteo manual antes de entregar"),
            ("Delta P(m) - P(m-1) -> A05 N/O", PENDIENTE,
             "Fase 4 del plan, no implementada", "Seguir usando el CALCULADOR A05"),
        ],
    },
    "sm_rescate_inasistentes": {
        "rem": "Auditoria (no tributa a ninguna casilla del REM)",
        "cubre": [],
        "no_cubre": [
            ("(todas)", FUERA,
             "Este reporte no tributa a ninguna casilla del REM: es auditoria "
             "operativa de rescate (SP_P6_poblacion_plan.md §8)", "No copiar nada de aqui al SP/SA"),
            ("Posibles_Fallecidos / Posibles_Traslados", VALIDACION,
             "El match de 'fallecido' / 'traslado-cambio de domicilio' es un heuristico "
             "sobre Motivo Pasivacion sin validar contra los valores reales del export. "
             "NINGUNO se excluye de Rescate_6m/13m (corregido sep-2026): quedan flageados",
             "Revisar que las listas capturen los Motivo Pasivacion de verdad"),
            ("Brecha_Medico", VALIDACION,
             "Compara dos pasadas de poblacion.construir_poblacion (con/sin filtro "
             "medico); todavia sin validar contra un caso real conocido",
             "Confirmar a mano un par de RUN antes de usarlo para gestion"),
        ],
    },
    # Gemelo de a05_o_egresos: mismo REM, misma tabla de "no_cubre" (§3.4 del plan).
    "a05_n_ingresos": None,
}
COBERTURA["a05_n_ingresos"] = {
    "rem": COBERTURA["a05_o_egresos"]["rem"],
    "cubre": ["A05 Ingresos"],
    "no_cubre": COBERTURA["a05_o_egresos"]["no_cubre"],
}


_SIN_AVISOS = "Sin avisos: todas las fuentes opcionales estaban cargadas."
_ENCABEZADO_TABLA = ("Casilla", "Estado", "Motivo", "Que hacer")


def entradas(modulo_id, avisos=()):
    """Filas de la hoja: lo estructural del catalogo (union, si `modulo_id` es una
    lista de ids que comparten UN archivo de salida, p.ej. A05 egresos+ingresos)
    + los `avisos` de esta corrida. `avisos` = iterable de
    (casilla, estado, motivo, que_hacer). Devuelve {rem, cubre, no_cubre, avisos}."""
    ids = list(modulo_id) if isinstance(modulo_id, (list, tuple)) else [modulo_id]
    faltan = [i for i in ids if i not in COBERTURA]
    if faltan:
        raise KeyError(f"cobertura: modulo(s) sin catalogo: {', '.join(faltan)}")
    cubre, no_cubre, vistas = [], [], set()
    for mid in ids:
        cat = COBERTURA[mid]
        for c in cat["cubre"]:
            if c not in cubre:
                cubre.append(c)
        for fila in cat["no_cubre"]:
            if fila not in vistas:
                vistas.add(fila)
                no_cubre.append(fila)
    return {"rem": COBERTURA[ids[0]]["rem"], "cubre": cubre, "no_cubre": no_cubre,
            "avisos": list(avisos)}


def escribir_hoja(wb, modulo_id, contexto=None, avisos=(), nombre="LEEME"):
    """Crea la hoja `nombre` como PRIMERA del workbook openpyxl `wb` (sirve tanto a
    un libro recien creado por pandas.ExcelWriter -sin hojas todavia- como a uno
    que ya trae la hoja de datos original en el indice 0, camino A05/A03).
    `contexto` = dict opcional con 'mes' ((año,mes) o texto) y 'archivos' (lista de
    nombres). Devuelve la worksheet creada."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    contexto = contexto or {}
    d = entradas(modulo_id, avisos)
    ws = wb.create_sheet(nombre, 0)

    def fila(*vals):
        ws.append(list(vals))
        return ws[ws.max_row]

    fila(f"autoREM {VERSION} - QUE NO INCLUYE ESTA PLANILLA")[0].font = Font(bold=True, size=12)
    fila(f"REM: {d['rem']}")
    mes = contexto.get("mes")
    mes_txt = f"{mes[0]}-{mes[1]:02d}" if isinstance(mes, (tuple, list)) and len(mes) == 2 else (mes or "(no especificado)")
    generado = contexto.get("generado") or datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    fila(f"Mes reportado: {mes_txt}     Generado: {generado}")
    archivos = contexto.get("archivos")
    if archivos:
        fila("Entradas: " + " - ".join(str(a) for a in archivos))
    if d["cubre"]:
        fila("Cubre: " + " - ".join(d["cubre"]))
    fila("")

    fila("ESTA CORRIDA")[0].font = Font(bold=True)
    for c in fila(*_ENCABEZADO_TABLA):
        c.font = Font(bold=True)
    if d["avisos"]:
        for casilla, estado, motivo, que_hacer in d["avisos"]:
            fila(casilla, estado, motivo, que_hacer)
    else:
        fila(_SIN_AVISOS)
    fila("")

    fila("SIEMPRE A MANO / FUERA DE ALCANCE")[0].font = Font(bold=True)
    for c in fila(*_ENCABEZADO_TABLA):
        c.font = Font(bold=True)
    for casilla, categoria, motivo, que_hacer in d["no_cubre"]:
        fila(casilla, categoria, motivo, que_hacer)

    ws.freeze_panes = "A2"
    for i, w in enumerate([46, 16, 62, 46], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws
